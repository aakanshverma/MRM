"""Excel workbook: sheets, expected columns, formatting and charts."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from model_monitoring.excel_report import SHEET_NAMES, build_workbook
from model_monitoring.excel_style import (
    RISK_RAMP,
    risk_fill_for_rank,
    sanitise_sheet_name,
    sanitise_table_name,
)

from helpers import make_config, make_frame


@pytest.fixture(scope="module")
def workbook(tmp_path_factory):
    """Build the workbook once and inspect it from several tests."""
    from model_monitoring.logging_config import setup_logging
    from model_monitoring.main import run_pipeline

    setup_logging("ERROR", quiet=True)
    frame = make_frame(n=1_600)
    config = make_config()
    directory = tmp_path_factory.mktemp("report")
    config.excel.output_dir = str(directory)
    result = run_pipeline(
        config, frame=frame, write_excel=False, output=lambda *a, **k: None
    )
    path = build_workbook(result.report_inputs, directory / "report.xlsx")
    return load_workbook(path), path, result


# ---------------------------------------------------------------------------
# Sheets
# ---------------------------------------------------------------------------


def test_every_required_sheet_is_present(workbook):
    """Requirement: the workbook contains the expected sheets."""
    book, _, _ = workbook
    expected = [
        SHEET_NAMES["dashboard"], SHEET_NAMES["mer"], SHEET_NAMES["segments"],
        SHEET_NAMES["psi"], SHEET_NAMES["csi"], SHEET_NAMES["csi_bins"],
        SHEET_NAMES["migration"], SHEET_NAMES["data_quality"],
        SHEET_NAMES["performance"], SHEET_NAMES["config"], SHEET_NAMES["alerts"],
        SHEET_NAMES["log"], SHEET_NAMES["bins"],
    ]
    for name in expected:
        assert name in book.sheetnames, f"missing sheet: {name}"


def test_sheet_names_are_valid_for_excel(workbook):
    book, _, _ = workbook
    for name in book.sheetnames:
        assert len(name) <= 31
        assert not set(name) & set("[]:*?/\\")


def test_dashboard_is_the_first_visible_sheet(workbook):
    book, _, _ = workbook
    visible = [name for name in book.sheetnames if book[name].sheet_state == "visible"]
    assert visible[0] == SHEET_NAMES["dashboard"]


def test_chart_data_sheet_is_hidden(workbook):
    book, _, _ = workbook
    assert book["_ChartData"].sheet_state == "hidden"


# ---------------------------------------------------------------------------
# Expected columns
# ---------------------------------------------------------------------------


def _locate_header(sheet, expected_first: str, search_rows: int = 40):
    """Find a table's header row by its first column name.

    Detail sheets open with a title and a methodology note, so the header is
    never row 1; tests must anchor to it rather than assume a fixed offset.
    """
    for row in range(1, search_rows):
        if sheet.cell(row=row, column=1).value == expected_first:
            return row, [
                sheet.cell(row=row, column=col).value
                for col in range(1, sheet.max_column + 1)
            ]
    raise AssertionError(f"header starting with {expected_first!r} not found")


def _header(sheet, expected_first: str, search_rows: int = 40) -> list:
    return _locate_header(sheet, expected_first, search_rows)[1]


def _column_values(sheet, first_column: str, column: str, limit: int = 60) -> set:
    """Values of one column, read from the row after its header."""
    header_row, header = _locate_header(sheet, first_column)
    index = header.index(column) + 1
    last = min(sheet.max_row, header_row + limit)
    return {
        sheet.cell(row=row, column=index).value
        for row in range(header_row + 1, last + 1)
    }


def test_mer_sheet_has_the_required_columns(workbook):
    book, _, _ = workbook
    header = _header(book[SHEET_NAMES["mer"]], "MODEL")
    for column in (
        "MODEL", "QUARTER", "TOTAL_RECORDS", "UNIQUE_LANS", "UNIQUE_CUSTOMERS",
        "MISSING_SEGMENT_COUNT", "MISSING_SEGMENT_PCT", "INVALID_SEGMENT_COUNT",
        "INVALID_SEGMENT_PCT", "MOST_POPULATED_SEGMENT", "SEGMENT_CONCENTRATION_PCT",
        "POPULATION_CHANGE_VS_REFERENCE_PCT", "POPULATION_CHANGE_VS_PREV_QUARTER_PCT",
        "MISSING_OUTPUT_PCT", "DUPLICATE_LAN_COUNT", "REFERENCE_LABEL", "REFERENCE_MODE",
        "PSI", "PSI_SEVERITY", "PSI_QOQ", "GINI", "PERFORMANCE_STATUS",
    ):
        assert column in header, f"MER sheet missing {column}"


def test_segment_distribution_sheet_has_the_required_columns(workbook):
    book, _, _ = workbook
    header = _header(book[SHEET_NAMES["segments"]], "MODEL")
    for column in (
        "MODEL", "QUARTER", "SEGMENT", "RISK_RANK", "RISK_BAND", "CUSTOMER_COUNT",
        "POPULATION_PCT", "REFERENCE_PCT", "CHANGE_PP", "COUNT_CHANGE",
    ):
        assert column in header


def test_psi_sheet_has_the_required_columns(workbook):
    book, _, _ = workbook
    header = _header(book[SHEET_NAMES["psi"]], "MODEL")
    for column in (
        "MODEL", "QUARTER", "REFERENCE_LABEL", "REFERENCE_COUNT", "MONITORING_COUNT",
        "SEGMENT", "REFERENCE_PROPORTION", "MONITORING_PROPORTION",
        "PSI_CONTRIBUTION", "TOTAL_PSI", "SEVERITY",
    ):
        assert column in header


def test_csi_summary_sheet_has_the_required_columns(workbook):
    book, _, _ = workbook
    header = _header(book[SHEET_NAMES["csi"]], "MODEL")
    for column in (
        "MODEL", "QUARTER", "VARIABLE", "VARIABLE_TYPE", "CSI", "SEVERITY",
        "REFERENCE_COUNT", "MONITORING_COUNT", "N_BINS", "MISSING_PCT_REF",
        "MISSING_PCT_MON", "MISSINGNESS_CHANGE_PP", "TOP_CONTRIBUTING_BIN",
        "UNSEEN_CATEGORY_COUNT", "RECOMMENDED_ACTION", "REFERENCE_LABEL",
    ):
        assert column in header


def test_csi_bin_sheet_has_the_required_columns(workbook):
    book, _, _ = workbook
    header = _header(book[SHEET_NAMES["csi_bins"]], "MODEL")
    for column in (
        "MODEL", "QUARTER", "VARIABLE", "BIN", "REFERENCE_COUNT", "MONITORING_COUNT",
        "REFERENCE_PROPORTION", "MONITORING_PROPORTION", "CSI_CONTRIBUTION", "TOTAL_CSI",
    ):
        assert column in header


def test_migration_sheet_has_the_required_columns(workbook):
    book, _, _ = workbook
    header = _header(book[SHEET_NAMES["migration"]], "MODEL")
    for column in (
        "MODEL", "COMPARISON", "MATCHED_CUSTOMERS", "UNCHANGED", "UPGRADES",
        "DOWNGRADES", "UPGRADE_PCT", "DOWNGRADE_PCT", "NET_DOWNGRADE_PP",
        "AVERAGE_RANK_MOVEMENT", "LARGEST_FLOW", "SEVERITY",
    ):
        assert column in header


def test_data_quality_sheet_has_the_required_columns(workbook):
    book, _, _ = workbook
    header = _header(book[SHEET_NAMES["data_quality"]], "CHECK_ID")
    for column in ("CHECK_ID", "CATEGORY", "CHECK", "SEVERITY", "DETAIL",
                   "RECOMMENDED_ACTION"):
        assert column in header


def test_alert_register_sheet_has_the_required_columns(workbook):
    book, _, _ = workbook
    header = _header(book[SHEET_NAMES["alerts"]], "ALERT_ID")
    for column in (
        "MODEL", "QUARTER", "METRIC", "VARIABLE_OR_SEGMENT", "METRIC_VALUE",
        "THRESHOLD", "SEVERITY", "REFERENCE_PERIOD", "MONITORING_PERIOD",
        "EXPLANATION", "RECOMMENDED_ACTION", "INVESTIGATION_STATUS",
    ):
        assert column in header


# ---------------------------------------------------------------------------
# Dashboard content
# ---------------------------------------------------------------------------


def _all_text(sheet, max_row=120, max_col=20) -> str:
    values = []
    for row in range(1, min(sheet.max_row, max_row) + 1):
        for col in range(1, min(sheet.max_column, max_col) + 1):
            value = sheet.cell(row=row, column=col).value
            if value is not None:
                values.append(str(value))
    return "\n".join(values)


def test_dashboard_reports_the_required_headline_items(workbook):
    book, _, result = workbook
    text = _all_text(book[SHEET_NAMES["dashboard"]])
    for label in (
        "Reporting date", "Scorecards monitored", "Quarters analysed",
        "Overall monitoring status", "RED alerts", "AMBER alerts", "GREY alerts",
        "GREEN checks", "Material data-quality issues",
        "TOP FIVE STABILITY CONCERNS", "RECOMMENDED PRIORITY INVESTIGATIONS",
        "REFERENCE PERIOD AND MODE PER SCORECARD",
    ):
        assert label in text, f"dashboard missing {label!r}"
    # CSI count labels carry the configured thresholds
    assert "Characteristics with CSI >=" in text


def test_dashboard_distinguishes_fixed_from_quarter_on_quarter_psi(workbook):
    book, _, _ = workbook
    text = _all_text(book[SHEET_NAMES["dashboard"]])
    assert "PSI vs reference" in text
    assert "PSI quarter-on-quarter" in text
    assert "Fixed-reference PSI compares the quarter" in text


def test_dashboard_shows_each_models_own_reference_period(workbook):
    book, _, result = workbook
    text = _all_text(book[SHEET_NAMES["dashboard"]])
    for model in result.config.active_models:
        assert model.reference.label in text
        assert model.reference.mode.value in text


def test_dashboard_has_the_required_charts(workbook):
    book, _, _ = workbook
    sheet = book[SHEET_NAMES["dashboard"]]
    charts = getattr(sheet, "_charts", [])
    assert len(charts) >= 5
    titles = []
    for chart in charts:
        try:
            titles.append(
                "".join(
                    run.t or "" for para in chart.title.tx.rich.p
                    for run in (para.r or [])
                )
            )
        except Exception:
            titles.append("")
    joined = " ".join(titles)
    assert "Fixed-reference PSI" in joined
    assert "Quarter-on-quarter PSI" in joined
    assert "riskiest grade" in joined or "Riskier" in joined
    assert "Alert count by severity" in joined
    assert "segment distribution by quarter" in joined


# ---------------------------------------------------------------------------
# Methodology sheet
# ---------------------------------------------------------------------------


def test_methodology_sheet_documents_formulas_thresholds_and_limitations(workbook):
    book, _, _ = workbook
    text = _all_text(book[SHEET_NAMES["config"]], max_row=250)
    for expected in (
        "EXECUTION METADATA", "REFERENCE PERIOD AND MODE PER SCORECARD",
        "SCORECARD AND COLUMN MAPPING", "SEGMENT ORDER AND RISK DIRECTION",
        "THRESHOLD CONFIGURATION", "BINNING RULES",
        "METRIC DEFINITIONS AND FORMULAS", "MONITORING PERIODS",
        "ASSUMPTIONS AND LIMITATIONS", "INPUT PROVENANCE AND VALIDATION",
    ):
        assert expected in text, f"methodology sheet missing {expected!r}"
    # the actual formulas must be present, not just their names
    assert "PSI = SUM_i (A_i - E_i) * ln(A_i / E_i)" in text
    assert "CSI_j = SUM_i (A_ij - E_ij) * ln(A_ij / E_ij)" in text
    assert "Gini = 2 * AUC - 1" in text
    # smoothing must be stated with its parameter
    assert "pseudocount" in text.lower()
    assert "A_TO_E_INCREASING" in text


def test_methodology_sheet_states_that_thresholds_are_conventions(workbook):
    book, _, _ = workbook
    text = _all_text(book[SHEET_NAMES["config"]], max_row=250)
    assert "not universal regulatory limits" in text
    assert "does not establish regulatory approval" in text


# ---------------------------------------------------------------------------
# Formatting
# ---------------------------------------------------------------------------


def test_detail_sheets_freeze_panes_and_filter(workbook):
    book, _, _ = workbook
    for key in ("mer", "segments", "psi", "csi", "csi_bins", "data_quality", "alerts"):
        sheet = book[SHEET_NAMES[key]]
        assert sheet.freeze_panes is not None, f"{key} has no frozen panes"
        has_filter = sheet.auto_filter.ref is not None or len(list(sheet.tables)) > 0
        assert has_filter, f"{key} has neither an autofilter nor an Excel table"


def test_conditional_formatting_is_applied(workbook):
    book, _, _ = workbook
    for key in ("mer", "psi", "csi", "data_quality", "alerts"):
        sheet = book[SHEET_NAMES[key]]
        assert len(sheet.conditional_formatting._cf_rules) > 0, f"{key} has no formatting"


def test_severity_is_always_available_as_text_not_only_colour(workbook):
    """Requirement: never rely on colour alone."""
    book, _, _ = workbook
    sheet = book[SHEET_NAMES["alerts"]]
    values = _column_values(sheet, "ALERT_ID", "SEVERITY", limit=200)
    values.discard(None)
    assert values <= {"GREEN", "AMBER", "RED", "GREY"}
    assert values


def test_risk_band_text_accompanies_the_colour_ramp(workbook):
    book, _, _ = workbook
    sheet = book[SHEET_NAMES["segments"]]
    values = _column_values(sheet, "MODEL", "RISK_BAND", limit=100)
    values.discard(None)
    assert values <= {
        "1-LOWEST RISK", "2-LOW RISK", "3-MEDIUM RISK", "4-HIGH RISK", "5-HIGHEST RISK",
    }


def test_number_formats_are_applied_to_index_columns(workbook):
    book, _, _ = workbook
    sheet = book[SHEET_NAMES["psi"]]
    header_row, header = _locate_header(sheet, "MODEL")
    index = header.index("TOTAL_PSI") + 1
    formats = {
        sheet.cell(row=row, column=index).number_format
        for row in range(header_row + 1, header_row + 30)
    }
    assert formats == {"0.0000"}


def test_excel_tables_are_created_with_unique_names(workbook):
    book, _, _ = workbook
    names = []
    for name in book.sheetnames:
        names.extend(list(book[name].tables))
    assert len(names) == len(set(names))
    assert len(names) >= 5


# ---------------------------------------------------------------------------
# Style helpers
# ---------------------------------------------------------------------------


def test_risk_palette_follows_the_rank_not_the_label():
    """Reversing the risk direction must reverse the colours."""
    safest = risk_fill_for_rank(1, 6)
    riskiest = risk_fill_for_rank(6, 6)
    assert safest.start_color.rgb.endswith(RISK_RAMP[0])
    assert riskiest.start_color.rgb.endswith(RISK_RAMP[-1])
    assert risk_fill_for_rank(None, 6) is None


def test_name_sanitisation():
    assert sanitise_sheet_name("a" * 50) == "a" * 31
    assert "[" not in sanitise_sheet_name("bad[name]")
    assert sanitise_table_name("1 bad-name!") .startswith("t_")
    assert " " not in sanitise_table_name("has spaces")


# ---------------------------------------------------------------------------
# Robustness
# ---------------------------------------------------------------------------


def test_workbook_builds_without_an_outcome_column(tmp_path):
    from model_monitoring.logging_config import setup_logging
    from model_monitoring.main import run_pipeline

    setup_logging("ERROR", quiet=True)
    frame = make_frame(n=900, with_outcome=False)
    config = make_config(with_outcome=False)
    config.excel.output_dir = str(tmp_path)
    result = run_pipeline(
        config, frame=frame, write_excel=False, output=lambda *a, **k: None
    )
    path = build_workbook(result.report_inputs, tmp_path / "no_outcome.xlsx")
    book = load_workbook(path)
    sheet = book[SHEET_NAMES["performance"]]
    text = _all_text(sheet, max_row=60)
    assert "NOT AVAILABLE" in text
    assert "METRIC AVAILABILITY" in text
    assert "outcome.outcome_column" in text


def test_workbook_builds_with_charts_disabled(tmp_path):
    from model_monitoring.logging_config import setup_logging
    from model_monitoring.main import run_pipeline

    setup_logging("ERROR", quiet=True)
    frame = make_frame(n=600)
    config = make_config()
    config.excel.include_charts = False
    config.excel.output_dir = str(tmp_path)
    result = run_pipeline(
        config, frame=frame, write_excel=False, output=lambda *a, **k: None
    )
    path = build_workbook(result.report_inputs, tmp_path / "no_charts.xlsx")
    book = load_workbook(path)
    assert len(getattr(book[SHEET_NAMES["dashboard"]], "_charts", [])) == 0


def test_row_cap_truncates_with_an_explicit_banner_and_sidecar(tmp_path):
    from model_monitoring.logging_config import setup_logging
    from model_monitoring.main import run_pipeline

    setup_logging("ERROR", quiet=True)
    frame = make_frame(n=900)
    config = make_config()
    config.excel.output_dir = str(tmp_path)
    config.excel.max_rows_per_sheet = 20          # force truncation
    result = run_pipeline(
        config, frame=frame, write_excel=False, output=lambda *a, **k: None
    )
    path = build_workbook(result.report_inputs, tmp_path / "capped.xlsx")
    book = load_workbook(path)
    text = _all_text(book[SHEET_NAMES["csi_bins"]], max_row=40)
    assert "No data has been dropped from the analysis" in text
    # the full table must have been written beside the workbook
    assert list(Path(tmp_path).glob("csi_bins_*.csv"))
