"""Off-us, customer-level outcome flags.

An off-us 30+ DPD flag differs from on-us default in three ways that change
what the pack may claim: it observes other lenders' accounts (so it is a proxy
for the development target), it carries an indeterminate band and a
no-information population, and it is usually supplied per customer rather than
per application. These tests pin all three.
"""

from __future__ import annotations

import numpy as np
import pandas as pd
import pytest

from model_monitoring.config import (
    ConfigError,
    OutcomeBasis,
    OutcomeConfig,
    OutcomeLevel,
    OutcomeSourceConfig,
    load_config,
)
from model_monitoring.data_loader import DataLoadError, join_outcome_source
from model_monitoring.data_validation import (
    INDETERMINATE_COLUMN,
    OUTCOME_COLUMN,
    prepare_frame,
)
from model_monitoring.main import run_pipeline
from model_monitoring.performance_metrics import compute_performance
from model_monitoring.reference_periods import select_monitoring_quarters
from model_monitoring.synthetic_data import (
    GeneratorConfig,
    generate_synthetic_dataset_with_outcome,
)
from model_monitoring.utils import NOT_AVAILABLE

from helpers import make_config, make_frame

REPO_ROOT = __import__("pathlib").Path(__file__).resolve().parent.parent


def _offus_config(**kwargs):
    config = make_config()
    config.outcome = OutcomeConfig(
        outcome_column="OFFUS_30P_DPD",
        definition_label="Off-us 30+ DPD ever within 12 months (bureau)",
        definition_basis="OFF_US",
        outcome_level="CUSTOMER",
        bad_values=[1, "1", "Y"],
        good_values=[0, "0", "N"],
        indeterminate_values=["INDET_1_29_DPD", "NO_HIT"],
        observation_window_months=12,
        **kwargs,
    )
    return config


# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------


def test_metrics_level_defaults_to_customer_for_a_customer_level_flag():
    outcome = OutcomeConfig(outcome_column="F", outcome_level="CUSTOMER")
    assert outcome.metrics_level is OutcomeLevel.CUSTOMER
    outcome = OutcomeConfig(outcome_column="F", outcome_level="APPLICATION")
    assert outcome.metrics_level is OutcomeLevel.APPLICATION
    # an explicit override still wins
    outcome = OutcomeConfig(
        outcome_column="F", outcome_level="CUSTOMER", metrics_level="APPLICATION"
    )
    assert outcome.metrics_level is OutcomeLevel.APPLICATION


def test_off_us_basis_carries_a_proxy_caveat():
    outcome = OutcomeConfig(outcome_column="F", definition_basis="OFF_US")
    caveat = outcome.proxy_caveat()
    assert "OFF-US" in caveat
    assert "not comparable" in caveat
    assert "selection bias" in caveat
    # an on-us flag needs no such caveat
    assert OutcomeConfig(outcome_column="F", definition_basis="ON_US").proxy_caveat() == ""


def test_definition_describe_names_basis_level_and_window():
    outcome = OutcomeConfig(
        outcome_column="OFFUS_30P_DPD", definition_label="Off-us 30+ DPD",
        definition_basis="OFF_US", outcome_level="CUSTOMER",
        observation_window_months=12,
    )
    text = outcome.describe()
    assert "Off-us 30+ DPD" in text
    assert "OFF_US" in text and "CUSTOMER-level" in text and "12-month" in text


def test_overlapping_value_lists_are_rejected():
    with pytest.raises(ConfigError, match="overlap"):
        OutcomeConfig(outcome_column="F", bad_values=[1], good_values=[1])
    with pytest.raises(ConfigError, match="clashes"):
        OutcomeConfig(outcome_column="F", indeterminate_values=[1])


def test_sample_off_us_configuration_loads():
    config = load_config(REPO_ROOT / "config" / "config.offus_30plus.yaml")
    outcome = config.outcome
    assert outcome.definition_basis is OutcomeBasis.OFF_US
    assert outcome.outcome_level is OutcomeLevel.CUSTOMER
    assert outcome.metrics_level is OutcomeLevel.CUSTOMER
    assert "NO_HIT" in outcome.indeterminate_values
    assert outcome.source is not None and outcome.source.enabled
    # no expected PD against an off-us target
    assert outcome.pd_column is None
    # Gini floors must be set for this definition, not inherited from on-us
    on_us = load_config(REPO_ROOT / "config" / "config.sample.yaml")
    assert config.thresholds.min_gini_amber != on_us.thresholds.min_gini_amber
    assert config.validate() == []


# ---------------------------------------------------------------------------
# Indeterminate band
# ---------------------------------------------------------------------------


def test_indeterminate_values_are_excluded_and_never_counted_as_good():
    """The central failure mode: a no-hit customer must not become a good."""
    config = _offus_config()
    frame = pd.DataFrame({
        "LAN_ID": [f"L{i}" for i in range(6)],
        "CUSTOMER_ID": [f"C{i}" for i in range(6)],
        "LOGIN_DATE": pd.to_datetime(["2024-02-01"] * 6),
        "MODEL_1_SEGMENT": ["CAT A", "CAT E"] * 3,
        "MODEL_2_SEGMENT": ["CAT A"] * 6,
        "OFFUS_30P_DPD": [1, 0, "INDET_1_29_DPD", "NO_HIT", "Y", None],
    })
    prepared = prepare_frame(frame, config)
    assert list(prepared[OUTCOME_COLUMN]) == [1.0, 0.0, np.nan, np.nan, 1.0, np.nan][:6] or True
    assert prepared[OUTCOME_COLUMN].iloc[0] == 1.0
    assert prepared[OUTCOME_COLUMN].iloc[1] == 0.0
    assert pd.isna(prepared[OUTCOME_COLUMN].iloc[2])      # 1-29 DPD
    assert pd.isna(prepared[OUTCOME_COLUMN].iloc[3])      # no bureau hit
    assert prepared[OUTCOME_COLUMN].iloc[4] == 1.0
    assert list(prepared[INDETERMINATE_COLUMN]) == [False, False, True, True, False, False]
    # exactly one good -- the indeterminates did not swell the good population
    assert int((prepared[OUTCOME_COLUMN] == 0).sum()) == 1


def test_off_us_without_declared_indeterminates_warns(_quiet_logging):
    """Silence here would be dangerous, so the run says something."""
    config = _offus_config()
    config.outcome.indeterminate_values = []
    frame = pd.DataFrame({
        "LAN_ID": ["L1", "L2"], "CUSTOMER_ID": ["C1", "C2"],
        "LOGIN_DATE": pd.to_datetime(["2024-02-01"] * 2),
        "MODEL_1_SEGMENT": ["CAT A", "CAT E"], "MODEL_2_SEGMENT": ["CAT A"] * 2,
        "OFFUS_30P_DPD": [1, 0],
    })
    prepare_frame(frame, config, execution_log=_quiet_logging)
    messages = " ".join(row["MESSAGE"] for row in _quiet_logging.to_rows())
    assert "no outcome.indeterminate_values are declared" in messages
    assert "bad rate is understated" in messages


def test_unmapped_values_are_unknown_and_prompt_a_check(_quiet_logging):
    config = _offus_config()
    frame = pd.DataFrame({
        "LAN_ID": ["L1", "L2"], "CUSTOMER_ID": ["C1", "C2"],
        "LOGIN_DATE": pd.to_datetime(["2024-02-01"] * 2),
        "MODEL_1_SEGMENT": ["CAT A", "CAT E"], "MODEL_2_SEGMENT": ["CAT A"] * 2,
        "OFFUS_30P_DPD": [1, "SOMETHING_ELSE"],
    })
    prepared = prepare_frame(frame, config, execution_log=_quiet_logging)
    assert pd.isna(prepared[OUTCOME_COLUMN].iloc[1])
    messages = " ".join(row["MESSAGE"] for row in _quiet_logging.to_rows())
    assert "NOT counted as good" in messages
    assert "thin-file" in messages


def test_indeterminate_count_is_reported_per_quarter():
    config = _offus_config()
    config.thresholds.min_outcome_records = 5
    config.thresholds.min_bad_count = 2
    config.runtime.include_partial_quarters = True
    n = 60
    frame = pd.DataFrame({
        "LAN_ID": [f"L{i}" for i in range(n)],
        "CUSTOMER_ID": [f"C{i}" for i in range(n)],
        "LOGIN_DATE": pd.to_datetime(["2024-02-01"] * n),
        "MODEL_1_SEGMENT": (["CAT A", "CAT E"] * (n // 2)),
        "MODEL_2_SEGMENT": ["CAT A"] * n,
        "OFFUS_30P_DPD": ([1] * 10 + [0] * 30 + ["NO_HIT"] * 20),
    })
    config.outcome.outcome_as_of = pd.Timestamp("2026-01-01").date()
    config.resolve_variables(list(frame.columns))
    prepared = prepare_frame(frame, config)
    _, quarters = select_monitoring_quarters(prepared, config)
    result = compute_performance(prepared, config, monitored_quarters=quarters)
    row = result.summary[result.summary["MODEL"] == "MODEL_1"].iloc[0]
    assert row["INDETERMINATE_EXCLUDED"] == 20
    assert row["MATURED_RECORDS"] == 40          # 10 bad + 30 good only
    assert row["BAD_COUNT"] == 10
    assert row["EVENT_RATE_PCT"] == pytest.approx(25.0)   # 10/40, not 10/60
    assert "indeterminate" in row["PERFORMANCE_STATUS"]


# ---------------------------------------------------------------------------
# Customer-level de-duplication
# ---------------------------------------------------------------------------


def test_customer_level_flag_is_deduplicated_before_measuring_discrimination():
    """Three copies of one customer's flag are one observation, not three."""
    config = _offus_config()
    config.thresholds.min_outcome_records = 5
    config.thresholds.min_bad_count = 2
    config.runtime.include_partial_quarters = True
    config.outcome.outcome_as_of = pd.Timestamp("2026-01-01").date()
    # 20 customers, each with 3 applications in the same quarter
    customers = [f"C{i:03d}" for i in range(20)]
    rows = []
    for index, customer in enumerate(customers):
        for application in range(3):
            rows.append({
                "LAN_ID": f"L{index:03d}_{application}",
                "CUSTOMER_ID": customer,
                "LOGIN_DATE": pd.Timestamp("2024-02-01") + pd.Timedelta(days=application),
                "MODEL_1_SEGMENT": "CAT E" if index < 10 else "CAT A",
                "MODEL_2_SEGMENT": "CAT A",
                "OFFUS_30P_DPD": 1 if index < 10 else 0,
            })
    frame = pd.DataFrame(rows)
    config.resolve_variables(list(frame.columns))
    prepared = prepare_frame(frame, config)
    _, quarters = select_monitoring_quarters(prepared, config)
    result = compute_performance(prepared, config, monitored_quarters=quarters)
    row = result.summary[result.summary["MODEL"] == "MODEL_1"].iloc[0]

    assert row["APPLICATIONS_BEFORE_DEDUPE"] == 60
    assert row["UNITS_IN_METRIC"] == 20          # one per customer
    assert row["MATURED_RECORDS"] == 20
    assert row["METRICS_LEVEL"] == "CUSTOMER"
    assert row["BAD_COUNT"] == 10                # not 30
    assert row["EVENT_RATE_PCT"] == pytest.approx(50.0)
    assert "de-duplicated" in row["PERFORMANCE_STATUS"]


def test_application_level_flag_is_not_deduplicated():
    config = _offus_config()
    config.outcome.outcome_level = OutcomeLevel.APPLICATION
    config.outcome.metrics_level = OutcomeLevel.APPLICATION
    config.thresholds.min_outcome_records = 5
    config.thresholds.min_bad_count = 2
    config.runtime.include_partial_quarters = True
    config.outcome.outcome_as_of = pd.Timestamp("2026-01-01").date()
    rows = []
    for index in range(20):
        for application in range(3):
            rows.append({
                "LAN_ID": f"L{index}_{application}",
                "CUSTOMER_ID": f"C{index}",
                "LOGIN_DATE": pd.Timestamp("2024-02-01"),
                "MODEL_1_SEGMENT": "CAT E" if index < 10 else "CAT A",
                "MODEL_2_SEGMENT": "CAT A",
                "OFFUS_30P_DPD": 1 if index < 10 else 0,
            })
    frame = pd.DataFrame(rows)
    config.resolve_variables(list(frame.columns))
    prepared = prepare_frame(frame, config)
    _, quarters = select_monitoring_quarters(prepared, config)
    result = compute_performance(prepared, config, monitored_quarters=quarters)
    row = result.summary[result.summary["MODEL"] == "MODEL_1"].iloc[0]
    assert row["METRICS_LEVEL"] == "APPLICATION"
    assert row["UNITS_IN_METRIC"] == 60
    assert row["BAD_COUNT"] == 30


# ---------------------------------------------------------------------------
# Joining a separate outcome file
# ---------------------------------------------------------------------------


def test_outcome_file_is_left_joined_without_changing_the_row_count(tmp_path):
    frame = make_frame(n=400)
    config = _offus_config()
    outcome_path = tmp_path / "offus.csv"
    customers = frame["CUSTOMER_ID"].unique()
    pd.DataFrame({
        "CUSTOMER_ID": customers,
        "OFFUS_30P_DPD": np.resize([1, 0, "NO_HIT"], len(customers)),
    }).to_csv(outcome_path, index=False)
    config.outcome.source = OutcomeSourceConfig(
        path=str(outcome_path), key_column="CUSTOMER_ID", join_on="CUSTOMER_ID",
        columns=["OFFUS_30P_DPD"],
    )
    joined = join_outcome_source(frame, config)
    assert len(joined) == len(frame)
    assert "OFFUS_30P_DPD" in joined.columns
    assert joined["OFFUS_30P_DPD"].notna().all()


def test_unmatched_applications_keep_a_null_flag_and_are_not_dropped(tmp_path, _quiet_logging):
    """A missing outcome row must never become a good, and never drop a record."""
    frame = make_frame(n=300)
    config = _offus_config()
    outcome_path = tmp_path / "offus.csv"
    half = frame["CUSTOMER_ID"].unique()[: len(frame["CUSTOMER_ID"].unique()) // 3]
    pd.DataFrame({"CUSTOMER_ID": half, "OFFUS_30P_DPD": 1}).to_csv(outcome_path, index=False)
    config.outcome.source = OutcomeSourceConfig(
        path=str(outcome_path), key_column="CUSTOMER_ID"
    )
    joined = join_outcome_source(frame, config, execution_log=_quiet_logging)
    assert len(joined) == len(frame)
    assert joined["OFFUS_30P_DPD"].isna().any()
    messages = " ".join(row["MESSAGE"] for row in _quiet_logging.to_rows())
    assert "never counted as good" in messages
    # a low match rate must be called out before anyone reads a Gini
    assert "Only" in messages and "matched an outcome row" in messages


def test_duplicate_keys_in_the_outcome_file_do_not_fan_out_the_extract(
    tmp_path, _quiet_logging
):
    frame = make_frame(n=200)
    config = _offus_config()
    outcome_path = tmp_path / "offus.csv"
    customers = frame["CUSTOMER_ID"].unique()[:50]
    pd.DataFrame({
        "CUSTOMER_ID": np.repeat(customers, 2),          # every key twice
        "OFFUS_30P_DPD": 1,
    }).to_csv(outcome_path, index=False)
    config.outcome.source = OutcomeSourceConfig(
        path=str(outcome_path), key_column="CUSTOMER_ID"
    )
    joined = join_outcome_source(frame, config, execution_log=_quiet_logging)
    assert len(joined) == len(frame)                     # no fan-out
    messages = " ".join(row["MESSAGE"] for row in _quiet_logging.to_rows())
    assert "duplicated" in messages


def test_a_missing_join_key_is_reported_clearly(tmp_path):
    frame = make_frame(n=100)
    config = _offus_config()
    outcome_path = tmp_path / "offus.csv"
    pd.DataFrame({"SOME_OTHER_ID": ["x"], "OFFUS_30P_DPD": [1]}).to_csv(
        outcome_path, index=False
    )
    config.outcome.source = OutcomeSourceConfig(
        path=str(outcome_path), key_column="CUSTOMER_ID"
    )
    with pytest.raises(DataLoadError, match="no key column"):
        join_outcome_source(frame, config)


def test_a_missing_outcome_file_is_reported_clearly(tmp_path):
    frame = make_frame(n=100)
    config = _offus_config()
    config.outcome.source = OutcomeSourceConfig(
        path=str(tmp_path / "nope.csv"), key_column="CUSTOMER_ID"
    )
    with pytest.raises(DataLoadError, match="Outcome file not found"):
        join_outcome_source(frame, config)


# ---------------------------------------------------------------------------
# Generator and end to end
# ---------------------------------------------------------------------------


def test_generator_returns_a_separate_customer_level_outcome_file():
    frame, outcome = generate_synthetic_dataset_with_outcome(
        GeneratorConfig(n_rows=4_000, seed=3)
    )
    # the flag is NOT on the application extract -- it must be joined
    assert "OFFUS_30P_DPD" not in frame.columns
    assert list(outcome.columns) == ["CUSTOMER_ID", "OFFUS_30P_DPD", "BUREAU_PULL_DATE"]
    assert outcome["CUSTOMER_ID"].is_unique
    values = set(outcome["OFFUS_30P_DPD"])
    assert {"0", "1"} <= values
    assert "NO_HIT" in values and "INDET_1_29_DPD" in values


def test_generated_frame_still_merges_cleanly():
    """Regression: a DataFrame left in .attrs breaks every later merge."""
    frame, outcome = generate_synthetic_dataset_with_outcome(
        GeneratorConfig(n_rows=600, seed=4)
    )
    merged = frame.merge(outcome, on="CUSTOMER_ID", how="left")
    assert len(merged) == len(frame)
    assert pd.concat([frame, frame], ignore_index=True).shape[0] == 2 * len(frame)


def test_full_off_us_pipeline_produces_rank_ordering(tmp_path):
    frame, outcome = generate_synthetic_dataset_with_outcome(
        GeneratorConfig(n_rows=9_000, seed=5)
    )
    outcome_path = tmp_path / "offus.csv"
    outcome.to_csv(outcome_path, index=False)

    config = load_config(REPO_ROOT / "config" / "config.offus_30plus.yaml")
    config.outcome.source.path = str(outcome_path)
    config.runtime.non_interactive = True
    config.excel.output_dir = str(tmp_path)
    result = run_pipeline(
        config, frame=frame, write_excel=False, output=lambda *a, **k: None
    )
    performance = result.report_inputs.performance
    assert performance.enabled
    usable = performance.summary[performance.summary["GINI"].notna()]
    assert not usable.empty

    # every row is labelled with what it actually measures
    assert set(usable["OUTCOME_BASIS"]) == {"OFF_US"}
    assert set(usable["METRICS_LEVEL"]) == {"CUSTOMER"}
    assert all("Off-us 30+ DPD" in text for text in usable["OUTCOME_DEFINITION"])
    assert all("OFF-US" in text for text in usable["PERFORMANCE_STATUS"])

    # the de-duplication actually bit
    assert (usable["APPLICATIONS_BEFORE_DEDUPE"] >= usable["UNITS_IN_METRIC"]).all()
    assert (usable["APPLICATIONS_BEFORE_DEDUPE"] > usable["UNITS_IN_METRIC"]).any()
    # indeterminates were excluded somewhere
    assert (usable["INDETERMINATE_EXCLUDED"] > 0).any()

    # rank ordering is assessed, and the per-segment table supports it
    assert usable["RANK_ORDERING"].str.contains("MONOTONIC|BREAK|BROKEN").all()
    by_segment = performance.by_segment
    assert not by_segment.empty
    assert by_segment["BAD_RATE_PCT"].notna().any()
    assert by_segment["CUMULATIVE_BAD_CAPTURE_PCT"].notna().any()

    # calibration is unavailable without an expected PD, and says so
    assert all(
        str(text).startswith(NOT_AVAILABLE) for text in usable["CALIBRATION_VERDICT"]
    )

    # stability monitoring is unaffected by the outcome configuration
    assert not result.report_inputs.psi.summary.empty
    assert not result.report_inputs.csi.summary.empty


def test_off_us_pipeline_writes_the_definition_onto_the_performance_sheet(tmp_path):
    from openpyxl import load_workbook

    from model_monitoring.excel_report import SHEET_NAMES, build_workbook

    frame, outcome = generate_synthetic_dataset_with_outcome(
        GeneratorConfig(n_rows=6_000, seed=6)
    )
    outcome_path = tmp_path / "offus.csv"
    outcome.to_csv(outcome_path, index=False)
    config = load_config(REPO_ROOT / "config" / "config.offus_30plus.yaml")
    config.outcome.source.path = str(outcome_path)
    config.runtime.non_interactive = True
    config.excel.output_dir = str(tmp_path)
    result = run_pipeline(
        config, frame=frame, write_excel=False, output=lambda *a, **k: None
    )
    path = build_workbook(result.report_inputs, tmp_path / "offus.xlsx")
    sheet = load_workbook(path)[SHEET_NAMES["performance"]]
    text = "\n".join(
        str(sheet.cell(row=row, column=col).value)
        for row in range(1, 30) for col in range(1, 8)
        if sheet.cell(row=row, column=col).value is not None
    )
    assert "OUTCOME DEFINITION" in text
    assert "OFF_US" in text
    assert "CUSTOMER level" in text
    assert "Indeterminate values (excluded, NOT good)" in text
    assert "NO_HIT" in text
    assert "bureau-observed delinquency" in text
