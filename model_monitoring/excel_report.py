"""Excel workbook generation.

One workbook per execution, built from the DataFrames the analytical modules
produce.  This module contains no calculations: if a number appears here it was
computed upstream, which keeps the business logic testable without opening a
spreadsheet.

Sheets
------
=========================  ===========================================================
``01 Exec Dashboard``      head-line status, KPI block, top concerns, charts
``02 MER Summary``         one row per model-quarter with PSI, CSI and performance
``03 Segment Distribution``model-quarter-segment mix vs the reference period
``04 PSI Detail``          per-segment PSI contributions
``05 CSI Summary``         every model-quarter-characteristic combination
``06 CSI Bin Detail``      reference vs monitoring distribution per bin
``07 Segment Migration``   summary, 6x6 matrices and heatmaps
``08 Data Quality``        checks, exceptions and severities
``09 Outcome Performance`` populated when outcomes exist, else the missing dependencies
``10 Config & Methodology``mappings, reference periods, formulas, thresholds, limits
``11 Alert Register``      the unified alert table
``12 Execution Log``       warnings and exceptions raised during the run
``13 Bin Definitions``     the fitted reference bin boundaries (governance evidence)
``14 Test Scenarios``      planted-scenario manifest (synthetic runs only)
=========================  ===========================================================
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Mapping, Optional, Sequence, Tuple

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from . import __version__
from .alert_engine import alert_counts, overall_status, priority_investigations, top_concerns
from .config import MonitoringConfig, Severity, SmoothingMethod
from .csi_calculator import CsiResult
from .data_loader import LoadResult
from .data_validation import ValidationReport
from .excel_style import (
    BAND,
    FORMAT_COUNT,
    FORMAT_INDEX,
    FORMAT_INDEX_SHORT,
    FORMAT_PCT,
    FORMAT_PP_SIGNED,
    FORMAT_PROPORTION,
    FORMAT_RATIO,
    NAVY,
    apply_color_scale,
    apply_data_bars,
    apply_diverging_scale,
    apply_heatmap,
    apply_severity_rules,
    apply_threshold_rules,
    autosize_columns,
    block_ref,
    column_ref,
    finalise,
    risk_fill_for_rank,
    sanitise_sheet_name,
    severity_fill,
    severity_font,
    write_frame,
    write_key_values,
    write_note,
    write_section,
    write_title,
    write_truncation_banner,
)
from .logging_config import ExecutionLog, get_logger
from .mer_metrics import MODEL_SUMMARY_COLUMNS
from .performance_metrics import PerformanceResult
from .psi_calculator import VS_PREVIOUS_QUARTER, VS_REFERENCE, PsiResult
from .reference_periods import ReferenceResolution, reference_summary_frame
from .segment_migration import COMPARISON_QOQ, COMPARISON_REFERENCE, MIGRATION_CAVEAT, migration_pivot
from .utils import NOT_AVAILABLE, QuarterInfo

logger = get_logger("excel_report")

CHART_DATA_SHEET = "_ChartData"

SHEET_NAMES = {
    "dashboard": "01 Exec Dashboard",
    "mer": "02 MER Summary",
    "segments": "03 Segment Distribution",
    "psi": "04 PSI Detail",
    "csi": "05 CSI Summary",
    "csi_bins": "06 CSI Bin Detail",
    "migration": "07 Segment Migration",
    "data_quality": "08 Data Quality",
    "performance": "09 Outcome Performance",
    "config": "10 Config & Methodology",
    "alerts": "11 Alert Register",
    "log": "12 Execution Log",
    "bins": "13 Bin Definitions",
    "scenarios": "14 Test Scenarios",
}


# ---------------------------------------------------------------------------
# Inputs
# ---------------------------------------------------------------------------


@dataclass
class ReportInputs:
    """Everything the workbook needs, already computed."""

    config: MonitoringConfig
    load_result: LoadResult
    validation: ValidationReport
    resolutions: Mapping[str, ReferenceResolution]
    all_quarters: Sequence[QuarterInfo]
    monitored_quarters: Sequence[QuarterInfo]
    mer: pd.DataFrame
    model_summary: pd.DataFrame
    segment_distribution: pd.DataFrame
    psi: PsiResult
    csi: CsiResult
    migration_matrix: pd.DataFrame
    migration_summary: pd.DataFrame
    data_quality: pd.DataFrame
    performance: PerformanceResult
    alerts: pd.DataFrame
    execution_log: ExecutionLog
    scenarios: Optional[pd.DataFrame] = None

    @property
    def latest_quarter(self) -> Optional[str]:
        if not self.monitored_quarters:
            return None
        return self.monitored_quarters[-1].label


# ---------------------------------------------------------------------------
# Builder
# ---------------------------------------------------------------------------


class ExcelReportBuilder:
    """Renders :class:`ReportInputs` into a formatted workbook."""

    def __init__(self, inputs: ReportInputs) -> None:
        self.inputs = inputs
        self.config = inputs.config
        self.excel = inputs.config.excel
        self.segments = inputs.config.segments
        self.workbook = Workbook()
        self.chart_data: Optional[Worksheet] = None
        self._chart_row = 1
        self.sidecars: List[str] = []

    # -- public ------------------------------------------------------------
    def build(self, path: str | Path) -> Path:
        output = Path(path)
        output.parent.mkdir(parents=True, exist_ok=True)

        self.workbook.remove(self.workbook.active)
        self.chart_data = self.workbook.create_sheet(CHART_DATA_SHEET)

        dashboard = self.workbook.create_sheet(SHEET_NAMES["dashboard"])
        self._build_mer_sheet()
        self._build_segment_sheet()
        self._build_psi_sheet()
        self._build_csi_summary_sheet()
        self._build_csi_bin_sheet()
        self._build_migration_sheet()
        self._build_data_quality_sheet()
        self._build_performance_sheet()
        self._build_config_sheet()
        self._build_alert_sheet()
        self._build_log_sheet()
        self._build_bin_definition_sheet()
        if self.inputs.scenarios is not None and not self.inputs.scenarios.empty:
            self._build_scenario_sheet()
        # the dashboard is written last so its chart data blocks already exist
        self._build_dashboard(dashboard)

        self.chart_data.sheet_state = "hidden"
        self.workbook.active = 0
        self.workbook.save(output)
        logger.info("Excel workbook written to %s", output)
        return output

    # -- helpers -----------------------------------------------------------
    def _sheet(self, key: str) -> Worksheet:
        return self.workbook.create_sheet(sanitise_sheet_name(SHEET_NAMES[key]))

    def _sidecar(self, frame: pd.DataFrame, name: str) -> Optional[str]:
        """Write the full table beside the workbook when a sheet must truncate."""
        try:
            directory = Path(self.excel.output_dir)
            directory.mkdir(parents=True, exist_ok=True)
            stamp = self.config.execution_timestamp or datetime.now()
            path = directory / f"{name}_{stamp:%Y%m%d_%H%M%S}.csv"
            frame.to_csv(path, index=False)
            self.sidecars.append(str(path))
            return str(path)
        except Exception as exc:  # pragma: no cover - never fail the pack on this
            logger.warning("Could not write the %s overflow CSV: %s", name, exc)
            return None

    def _write_detail(
        self,
        key: str,
        frame: pd.DataFrame,
        *,
        title: str,
        subtitle: str,
        note: Optional[str] = None,
        number_formats: Optional[Mapping[str, str]] = None,
        table_name: Optional[str] = None,
        freeze_columns: int = 2,
        width_overrides: Optional[Mapping[str, int]] = None,
        empty_message: str = "No rows were produced for this sheet.",
    ) -> Tuple[Worksheet, Dict[str, Any]]:
        """Standard detail sheet: title, optional note, formatted table."""
        sheet = self._sheet(key)
        width = max(4, min(len(frame.columns), 16))
        row = write_title(sheet, title, subtitle, width=width)
        if note:
            row = write_note(sheet, row, note, width=width)
            row += 1
        if frame is None or frame.empty:
            write_note(sheet, row, empty_message, width=width)
            return sheet, {"rows_written": 0, "columns": {}}

        layout = write_frame(
            sheet, frame, start_row=row,
            number_formats=number_formats,
            max_rows=self.excel.max_rows_per_sheet,
            style_row_limit=self.excel.conditional_formatting_row_limit,
        )
        sidecar = (
            self._sidecar(frame, key) if layout["truncated"] else None
        )
        if layout["truncated"]:
            write_truncation_banner(sheet, layout["last_row"] + 2, layout, sidecar, width=width)
        autosize_columns(sheet, frame, overrides=width_overrides)
        finalise(
            sheet, layout,
            freeze=self.excel.freeze_panes,
            autofilter=self.excel.autofilter,
            table_name=table_name if self.excel.use_excel_tables else None,
            table_row_limit=self.excel.excel_table_row_limit,
            freeze_columns=freeze_columns,
        )
        return sheet, layout

    def _can_format(self, layout: Mapping[str, Any]) -> bool:
        return 0 < layout.get("rows_written", 0) <= self.excel.conditional_formatting_row_limit

    # -- chart data --------------------------------------------------------
    def _write_chart_block(self, title: str, frame: pd.DataFrame) -> Optional[Dict[str, Any]]:
        """Write a pivot block onto the hidden chart-data sheet."""
        if self.chart_data is None or frame is None or frame.empty:
            return None
        sheet = self.chart_data
        start = self._chart_row
        sheet.cell(row=start, column=1, value=title)
        header_row = start + 1
        sheet.cell(row=header_row, column=1, value=str(frame.index.name or "CATEGORY"))
        for offset, column in enumerate(frame.columns, start=2):
            sheet.cell(row=header_row, column=offset, value=str(column))
        for row_offset, (index_value, series) in enumerate(frame.iterrows(), start=1):
            sheet.cell(row=header_row + row_offset, column=1, value=str(index_value))
            for col_offset, value in enumerate(series.tolist(), start=2):
                if value is None or (isinstance(value, float) and value != value):
                    value = None
                elif hasattr(value, "item"):
                    value = value.item()
                sheet.cell(row=header_row + row_offset, column=col_offset, value=value)
        last_row = header_row + len(frame)
        self._chart_row = last_row + 2
        return {
            "header_row": header_row,
            "first_data_row": header_row + 1,
            "last_row": last_row,
            "first_col": 2,
            "last_col": 1 + len(frame.columns),
            "n_series": len(frame.columns),
        }

    def _line_chart(
        self, block: Optional[Dict[str, Any]], title: str, y_title: str, x_title: str = "Quarter",
    ) -> Optional[LineChart]:
        if not block:
            return None
        chart = LineChart()
        chart.title = title
        chart.style = 12
        chart.y_axis.title = y_title
        chart.x_axis.title = x_title
        chart.height = 7.5
        chart.width = 16
        chart.add_data(
            Reference(
                self.chart_data, min_col=block["first_col"], max_col=block["last_col"],
                min_row=block["header_row"], max_row=block["last_row"],
            ),
            titles_from_data=True,
        )
        chart.set_categories(
            Reference(
                self.chart_data, min_col=1,
                min_row=block["first_data_row"], max_row=block["last_row"],
            )
        )
        for series in chart.series:
            series.smooth = False
        return chart

    def _bar_chart(
        self,
        block: Optional[Dict[str, Any]],
        title: str,
        y_title: str,
        *,
        stacked: bool = False,
        horizontal: bool = False,
        x_title: str = "Quarter",
        show_values: bool = False,
    ) -> Optional[BarChart]:
        if not block:
            return None
        chart = BarChart()
        chart.type = "bar" if horizontal else "col"
        chart.title = title
        chart.style = 10
        chart.y_axis.title = y_title
        chart.x_axis.title = x_title
        chart.height = 7.5
        chart.width = 16
        if stacked:
            chart.grouping = "percentStacked"
            chart.overlap = 100
        chart.add_data(
            Reference(
                self.chart_data, min_col=block["first_col"], max_col=block["last_col"],
                min_row=block["header_row"], max_row=block["last_row"],
            ),
            titles_from_data=True,
        )
        chart.set_categories(
            Reference(
                self.chart_data, min_col=1,
                min_row=block["first_data_row"], max_row=block["last_row"],
            )
        )
        if show_values:
            chart.dataLabels = DataLabelList()
            chart.dataLabels.showVal = True
        return chart

    # ==================================================================
    # Sheet 1 -- Executive dashboard
    # ==================================================================
    def _build_dashboard(self, sheet: Worksheet) -> None:
        inputs = self.inputs
        config = self.config
        width = 14
        counts = alert_counts(inputs.alerts)
        status = overall_status(inputs.alerts)
        latest = inputs.latest_quarter
        stamp = config.execution_timestamp or datetime.now()

        row = write_title(
            sheet,
            config.report_title,
            f"{config.entity_name}  |  quarterly scorecard monitoring  |  framework v"
            f"{config.framework_version} (engine v{__version__})",
            width=width,
            note=(
                f"Generated {stamp:%Y-%m-%d %H:%M:%S} from "
                f"{inputs.load_result.source_path} ({inputs.load_result.rows_read:,} rows). "
                "PSI/CSI thresholds are conventional monitoring guidelines; a breach triggers "
                "investigation, not automatic model rejection."
            ),
        )

        # -- head-line status ------------------------------------------
        row = write_section(sheet, row, "MONITORING STATUS", width=width)
        sheet.merge_cells(start_row=row, start_column=1, end_row=row + 1, end_column=width)
        status_cell = sheet.cell(row=row, column=1, value=status)
        status_cell.font = Font(bold=True, size=16, color=NAVY)
        severity_key = status.split(" ")[0]
        fill = severity_fill(severity_key if severity_key in counts else "GREY")
        if fill is not None:
            status_cell.fill = fill
        status_cell.alignment = Alignment(horizontal="center", vertical="center")
        row += 3

        incomplete = [q.label for q in inputs.all_quarters if not q.is_complete]
        csi_amber_count = inputs.csi.count_above(config.thresholds.csi.amber, quarter=latest)
        csi_red_count = inputs.csi.count_above(config.thresholds.csi.red, quarter=latest)
        material_dq = (
            int(inputs.data_quality["SEVERITY"].isin(["RED", "AMBER"]).sum())
            if not inputs.data_quality.empty else 0
        )
        items: List[Tuple[str, Any]] = [
            ("Reporting date", f"{stamp:%Y-%m-%d}"),
            ("Latest monitored quarter", latest or NOT_AVAILABLE),
            ("Scorecards monitored", len(config.active_models)),
            ("Quarters analysed", len(inputs.monitored_quarters)),
            ("Quarters excluded as incomplete",
             ", ".join(incomplete) if incomplete else "none"),
            ("Overall monitoring status", status),
            ("RED alerts (escalate)", counts.get("RED", 0)),
            ("AMBER alerts (investigate)", counts.get("AMBER", 0)),
            ("GREY alerts (insufficient data)", counts.get("GREY", 0)),
            ("GREEN checks (within limits)", counts.get("GREEN", 0)),
            (f"Characteristics with CSI >= {config.thresholds.csi.amber:g} ({latest})",
             csi_amber_count),
            (f"Characteristics with CSI >= {config.thresholds.csi.red:g} ({latest})",
             csi_red_count),
            ("Material data-quality issues", material_dq),
            ("Records analysed", inputs.load_result.rows_read),
        ]
        severities = {
            "Overall monitoring status": severity_key,
            "RED alerts (escalate)": "RED" if counts.get("RED") else "GREEN",
            "AMBER alerts (investigate)": "AMBER" if counts.get("AMBER") else "GREEN",
            "GREY alerts (insufficient data)": "GREY" if counts.get("GREY") else "GREEN",
            "GREEN checks (within limits)": "GREEN",
            f"Characteristics with CSI >= {config.thresholds.csi.red:g} ({latest})":
                "RED" if csi_red_count else "GREEN",
            f"Characteristics with CSI >= {config.thresholds.csi.amber:g} ({latest})":
                "AMBER" if csi_amber_count else "GREEN",
            "Material data-quality issues": "AMBER" if material_dq else "GREEN",
        }
        row = write_key_values(sheet, row, items, value_col=4, severities=severities, label_width=3)
        row += 1

        # -- per-scorecard snapshot -------------------------------------
        row = write_section(
            sheet, row, "SCORECARD SNAPSHOT - LATEST QUARTER", width=width,
            note=(
                "Fixed-reference PSI compares the quarter with each scorecard's OWN approved "
                "baseline (shown per row). Quarter-on-quarter PSI compares it with the "
                "immediately preceding quarter and is the earlier warning of recent change."
            ),
        )
        snapshot = self._dashboard_snapshot()
        if not snapshot.empty:
            layout = write_frame(
                sheet, snapshot, start_row=row,
                number_formats={
                    "PSI vs reference": FORMAT_INDEX,
                    "PSI quarter-on-quarter": FORMAT_INDEX,
                    "Worst CSI": FORMAT_INDEX,
                    "Riskier-grade share %": FORMAT_PCT,
                    "Riskier-share change pp": FORMAT_PP_SIGNED,
                    "Records in quarter": FORMAT_COUNT,
                    "Reference records": FORMAT_COUNT,
                },
            )
            for name in ("Status", "PSI severity", "QoQ severity"):
                apply_severity_rules(sheet, column_ref(layout, name))
            apply_threshold_rules(
                sheet, column_ref(layout, "PSI vs reference"),
                config.thresholds.psi.amber, config.thresholds.psi.red,
            )
            apply_threshold_rules(
                sheet, column_ref(layout, "PSI quarter-on-quarter"),
                config.thresholds.psi.amber, config.thresholds.psi.red,
            )
            apply_diverging_scale(sheet, column_ref(layout, "Riskier-share change pp"))
            autosize_columns(sheet, snapshot)
            row = layout["last_row"] + 2

        # -- top concerns ----------------------------------------------
        row = write_section(
            sheet, row, "TOP FIVE STABILITY CONCERNS", width=width,
            note="Ranked by severity then magnitude. Each row links to a detail sheet below.",
        )
        concerns = top_concerns(inputs.alerts, n=5, latest_quarter=latest)
        if not concerns.empty:
            layout = write_frame(sheet, concerns, start_row=row)
            apply_severity_rules(sheet, column_ref(layout, "SEVERITY"))
            autosize_columns(
                sheet, concerns,
                overrides={"CONCERN": 90, "METRIC": 28, "VARIABLE_OR_SEGMENT": 24},
            )
            for target in range(layout["first_data_row"], layout["last_row"] + 1):
                sheet.cell(row=target, column=layout["columns"]["CONCERN"]).alignment = Alignment(
                    wrap_text=True, vertical="top"
                )
                sheet.row_dimensions[target].height = 46
            row = layout["last_row"] + 2

        # -- priority investigations ------------------------------------
        row = write_section(
            sheet, row, "RECOMMENDED PRIORITY INVESTIGATIONS", width=width,
            note=(
                "De-duplicated by model and metric. Investigation status is carried in the "
                "Alert Register, where the model owner records findings and sign-off."
            ),
        )
        priorities = priority_investigations(inputs.alerts, n=10)
        if not priorities.empty:
            layout = write_frame(sheet, priorities, start_row=row)
            apply_severity_rules(sheet, column_ref(layout, "SEVERITY"))
            autosize_columns(
                sheet, priorities,
                overrides={"RECOMMENDED_ACTION": 90, "METRIC": 28, "VARIABLE_OR_SEGMENT": 24},
            )
            for target in range(layout["first_data_row"], layout["last_row"] + 1):
                sheet.cell(
                    row=target, column=layout["columns"]["RECOMMENDED_ACTION"]
                ).alignment = Alignment(wrap_text=True, vertical="top")
                sheet.row_dimensions[target].height = 46
            row = layout["last_row"] + 2

        # -- reference periods ------------------------------------------
        row = write_section(
            sheet, row, "REFERENCE PERIOD AND MODE PER SCORECARD", width=width,
            note=(
                "Each scorecard is monitored against its own baseline because the four models "
                "were developed, implemented and retrained at different times. No model's "
                "window is substituted for another's."
            ),
        )
        reference = reference_summary_frame(inputs.resolutions)
        if not reference.empty:
            display = reference[[
                "MODEL", "REFERENCE_MODE", "REFERENCE_LABEL", "RECORDS_IN_REFERENCE",
                "SCORED_RECORDS_IN_REFERENCE", "SHARE_OF_DATASET_PCT", "SUFFICIENT_SAMPLE",
                "CONFIGURED_BY", "WARNINGS",
            ]]
            layout = write_frame(
                sheet, display, start_row=row,
                number_formats={
                    "RECORDS_IN_REFERENCE": FORMAT_COUNT,
                    "SCORED_RECORDS_IN_REFERENCE": FORMAT_COUNT,
                    "SHARE_OF_DATASET_PCT": FORMAT_PCT,
                },
            )
            autosize_columns(sheet, display, overrides={"WARNINGS": 70})
            row = layout["last_row"] + 2

        if self.excel.include_charts:
            self._add_dashboard_charts(sheet, row)

        sheet.sheet_view.showGridLines = False
        for index in range(1, width + 1):
            letter = get_column_letter(index)
            if sheet.column_dimensions[letter].width in (None, 0):
                sheet.column_dimensions[letter].width = 16

    def _dashboard_snapshot(self) -> pd.DataFrame:
        """One row per scorecard for the latest monitored quarter."""
        summary = self.inputs.model_summary
        if summary is None or summary.empty:
            return pd.DataFrame()
        rows = []
        for record in summary.to_dict("records"):
            rows.append({
                "Scorecard": record["MODEL"],
                "Latest quarter": record.get("LATEST_QUARTER"),
                "Records in quarter": record.get("LATEST_QUARTER_RECORDS"),
                "Reference period": record.get("REFERENCE_LABEL"),
                "Reference mode": record.get("REFERENCE_MODE"),
                "Reference records": record.get("REFERENCE_RECORDS"),
                "PSI vs reference": record.get("LATEST_PSI"),
                "PSI severity": record.get("LATEST_PSI_SEVERITY"),
                "PSI quarter-on-quarter": record.get("LATEST_QOQ_PSI"),
                "QoQ severity": record.get("LATEST_QOQ_PSI_SEVERITY"),
                "CSI AMBER": record.get("CSI_VARIABLES_AMBER"),
                "CSI RED": record.get("CSI_VARIABLES_RED"),
                "Worst CSI variable": record.get("WORST_CSI_VARIABLE"),
                "Worst CSI": record.get("WORST_CSI"),
                "Riskier-grade share %": record.get("LATEST_RISKIER_SHARE_PCT"),
                "Riskier-share change pp": record.get("RISKIER_SHARE_CHANGE_PP"),
                "Gini": record.get("LATEST_GINI"),
                "Status": record.get("OVERALL_STATUS"),
            })
        return pd.DataFrame(rows)

    def _add_dashboard_charts(self, sheet: Worksheet, start_row: int) -> None:
        """Charts for PSI trends, segment mix, riskier share, alerts and CSI."""
        inputs = self.inputs
        row = write_section(sheet, start_row, "CHARTS", width=14, note=(
            "Fixed-reference PSI and quarter-on-quarter PSI are charted separately: they answer "
            "different questions and must not be read as one series."
        ))
        anchor_row = row + 1

        psi_fixed = self._pivot_psi(VS_REFERENCE)
        block = self._write_chart_block("PSI vs reference by model", psi_fixed)
        chart = self._line_chart(
            block, "Fixed-reference PSI by scorecard (vs each model's own baseline)",
            "PSI (index)",
        )
        if chart is not None:
            sheet.add_chart(chart, f"A{anchor_row}")

        psi_qoq = self._pivot_psi(VS_PREVIOUS_QUARTER)
        block = self._write_chart_block("PSI quarter-on-quarter by model", psi_qoq)
        chart = self._line_chart(
            block, "Quarter-on-quarter PSI by scorecard (vs preceding quarter)", "PSI (index)"
        )
        if chart is not None:
            sheet.add_chart(chart, f"J{anchor_row}")
        anchor_row += 16

        riskier = self._pivot_mer("RISKIER_SEGMENT_SHARE_PCT")
        block = self._write_chart_block("Riskier-grade share by model", riskier)
        chart = self._line_chart(
            block,
            f"Share of applications in the riskiest grade(s) "
            f"({', '.join(self.segments.riskier_segments)})",
            "% of scored population",
        )
        if chart is not None:
            sheet.add_chart(chart, f"A{anchor_row}")

        counts = alert_counts(inputs.alerts)
        severity_frame = pd.DataFrame(
            {"Alerts": [counts.get(key, 0) for key in ("RED", "AMBER", "GREY", "GREEN")]},
            index=pd.Index(["RED", "AMBER", "GREY", "GREEN"], name="SEVERITY"),
        )
        block = self._write_chart_block("Alert count by severity", severity_frame)
        chart = self._bar_chart(
            block, "Alert count by severity", "Number of alerts", x_title="Severity",
            show_values=True,
        )
        if chart is not None:
            sheet.add_chart(chart, f"J{anchor_row}")
        anchor_row += 16

        top_csi = inputs.csi.top_variables(
            n=min(15, self.excel.top_n), quarter=inputs.latest_quarter
        )
        if not top_csi.empty:
            labels = (top_csi["MODEL"] + " | " + top_csi["VARIABLE"]).tolist()
            csi_frame = pd.DataFrame(
                {"CSI": top_csi["CSI"].astype("float64").tolist()},
                index=pd.Index(labels, name="CHARACTERISTIC"),
            )
            block = self._write_chart_block("Top CSI characteristics", csi_frame)
            chart = self._bar_chart(
                block, f"Highest-CSI characteristics ({inputs.latest_quarter})", "CSI (index)",
                horizontal=True, x_title="Characteristic", show_values=True,
            )
            if chart is not None:
                sheet.add_chart(chart, f"A{anchor_row}")

        for model in self.config.active_models:
            mix = self._pivot_segment_mix(model.name)
            block = self._write_chart_block(f"Segment mix {model.name}", mix)
            chart = self._bar_chart(
                block, f"{model.name}: segment distribution by quarter", "% of scored population",
                stacked=True,
            )
            if chart is not None:
                sheet.add_chart(chart, f"J{anchor_row}")
                anchor_row += 16

    def _pivot_psi(self, comparison: str) -> pd.DataFrame:
        summary = self.inputs.psi.summary
        if summary is None or summary.empty:
            return pd.DataFrame()
        subset = summary[summary["COMPARISON"] == comparison]
        if subset.empty:
            return pd.DataFrame()
        pivot = subset.pivot_table(index="QUARTER", columns="MODEL", values="PSI", aggfunc="first")
        pivot.index.name = "QUARTER"
        return pivot.sort_index()

    def _pivot_mer(self, column: str) -> pd.DataFrame:
        mer = self.inputs.mer
        if mer is None or mer.empty or column not in mer.columns:
            return pd.DataFrame()
        pivot = mer.pivot_table(index="QUARTER", columns="MODEL", values=column, aggfunc="first")
        pivot.index.name = "QUARTER"
        return pivot.sort_index()

    def _pivot_segment_mix(self, model: str) -> pd.DataFrame:
        distribution = self.inputs.segment_distribution
        if distribution is None or distribution.empty:
            return pd.DataFrame()
        subset = distribution[distribution["MODEL"] == model]
        if subset.empty:
            return pd.DataFrame()
        pivot = subset.pivot_table(
            index="QUARTER", columns="SEGMENT", values="POPULATION_PCT", aggfunc="first"
        )
        ordered = [s for s in self.segments.segments_by_risk if s in pivot.columns]
        pivot = pivot.reindex(columns=ordered)
        pivot.index.name = "QUARTER"
        return pivot.sort_index()

    # ==================================================================
    # Sheet 2 -- MER summary
    # ==================================================================
    def _build_mer_sheet(self) -> None:
        formats = {
            "TOTAL_RECORDS": FORMAT_COUNT, "UNIQUE_LANS": FORMAT_COUNT,
            "UNIQUE_CUSTOMERS": FORMAT_COUNT, "SCORED_RECORDS": FORMAT_COUNT,
            "REFERENCE_RECORDS": FORMAT_COUNT, "DUPLICATE_LAN_COUNT": FORMAT_COUNT,
            "MISSING_SEGMENT_COUNT": FORMAT_COUNT, "INVALID_SEGMENT_COUNT": FORMAT_COUNT,
            "PREVIOUS_QUARTER_RECORDS": FORMAT_COUNT,
            "MISSING_SEGMENT_PCT": FORMAT_PCT, "INVALID_SEGMENT_PCT": FORMAT_PCT,
            "MISSING_OUTPUT_PCT": FORMAT_PCT, "DUPLICATE_LAN_PCT": FORMAT_PCT,
            "MOST_POPULATED_SEGMENT_PCT": FORMAT_PCT, "SEGMENT_CONCENTRATION_PCT": FORMAT_PCT,
            "RISKIER_SEGMENT_SHARE_PCT": FORMAT_PCT, "RISKIER_SEGMENT_SHARE_REF_PCT": FORMAT_PCT,
            "RISKIER_SEGMENT_SHARE_CHANGE_PP": FORMAT_PP_SIGNED,
            "POPULATION_CHANGE_VS_REFERENCE_PCT": FORMAT_PCT,
            "POPULATION_CHANGE_VS_PREV_QUARTER_PCT": FORMAT_PCT,
            "SEGMENT_HHI": FORMAT_PROPORTION, "PSI": FORMAT_INDEX, "PSI_QOQ": FORMAT_INDEX,
            "WORST_CSI": FORMAT_INDEX, "AUC": FORMAT_INDEX, "GINI": FORMAT_INDEX,
            "KS_STATISTIC": FORMAT_INDEX, "EVENT_RATE_PCT": FORMAT_PCT,
            "NON_EVENT_RATE_PCT": FORMAT_PCT, "AVERAGE_SEGMENT_RANK": FORMAT_RATIO,
            "AVERAGE_SEGMENT_RANK_CHANGE": FORMAT_RATIO,
        }
        for segment in self.segments.segment_order:
            formats[f"PCT_{segment}"] = FORMAT_PCT
            formats[f"N_{segment}"] = FORMAT_COUNT

        sheet, layout = self._write_detail(
            "mer", self.inputs.mer,
            title="Model Effectiveness Report - quarterly summary",
            subtitle="One row per scorecard and monitoring quarter",
            note=(
                "Volume, coverage and segment-mix metrics, with PSI, CSI counts and (where "
                "outcome labels exist) performance metrics merged in. Model STABILITY and model "
                "PREDICTIVE PERFORMANCE are reported separately: segment output alone supports "
                "the former and says nothing about the latter. "
                "QUARTER_IN_REFERENCE_WINDOW=TRUE means the quarter overlaps its own baseline, "
                "so its fixed-reference PSI is mechanically near zero."
            ),
            number_formats=formats,
            table_name="tbl_mer_summary",
        )
        if not self._can_format(layout):
            return
        for name in ("PSI_SEVERITY", "PSI_QOQ_SEVERITY", "PERFORMANCE_SEVERITY"):
            apply_severity_rules(sheet, column_ref(layout, name))
        apply_threshold_rules(
            sheet, column_ref(layout, "PSI"),
            self.config.thresholds.psi.amber, self.config.thresholds.psi.red,
        )
        apply_threshold_rules(
            sheet, column_ref(layout, "PSI_QOQ"),
            self.config.thresholds.psi.amber, self.config.thresholds.psi.red,
        )
        apply_color_scale(sheet, column_ref(layout, "MISSING_OUTPUT_PCT"))
        apply_color_scale(sheet, column_ref(layout, "SEGMENT_CONCENTRATION_PCT"))
        apply_diverging_scale(sheet, column_ref(layout, "RISKIER_SEGMENT_SHARE_CHANGE_PP"))
        apply_diverging_scale(sheet, column_ref(layout, "POPULATION_CHANGE_VS_PREV_QUARTER_PCT"))
        apply_data_bars(sheet, column_ref(layout, "TOTAL_RECORDS"))
        mix_columns = [f"PCT_{segment}" for segment in self.segments.segment_order]
        apply_heatmap(sheet, block_ref(layout, mix_columns))

    # ==================================================================
    # Sheet 3 -- Segment distribution
    # ==================================================================
    def _build_segment_sheet(self) -> None:
        sheet, layout = self._write_detail(
            "segments", self.inputs.segment_distribution,
            title="Segment distribution by scorecard, quarter and risk grade",
            subtitle="Population mix against each scorecard's own reference period",
            note=(
                "RISK_BAND is the text equivalent of the colour ramp, so the sheet is readable "
                "in greyscale and to colour-blind readers. The palette follows the configured "
                f"risk direction ({self.segments.risk_direction.value}): "
                f"{self.segments.safest_segment} is the safest grade and "
                f"{self.segments.riskiest_segment} the riskiest. MATERIAL_SHIFT marks a change "
                "of 5 percentage points or more against the reference period."
            ),
            number_formats={
                "CUSTOMER_COUNT": FORMAT_COUNT, "REFERENCE_COUNT": FORMAT_COUNT,
                "COUNT_CHANGE": FORMAT_COUNT, "QUARTER_TOTAL": FORMAT_COUNT,
                "REFERENCE_TOTAL": FORMAT_COUNT,
                "POPULATION_PCT": FORMAT_PCT, "REFERENCE_PCT": FORMAT_PCT,
                "CHANGE_PP": FORMAT_PP_SIGNED, "RELATIVE_CHANGE_PCT": FORMAT_PCT,
            },
            table_name="tbl_segment_distribution",
            freeze_columns=3,
        )
        if not self._can_format(layout):
            return
        apply_diverging_scale(sheet, column_ref(layout, "CHANGE_PP"))
        apply_data_bars(sheet, column_ref(layout, "POPULATION_PCT"))
        apply_color_scale(sheet, column_ref(layout, "RISK_RANK"))
        # risk-ramp fill on the segment label, honouring the configured direction
        segment_col = layout["columns"].get("SEGMENT")
        rank_col = layout["columns"].get("RISK_RANK")
        if segment_col and rank_col:
            for target in range(layout["first_data_row"], layout["last_row"] + 1):
                rank = sheet.cell(row=target, column=rank_col).value
                fill = risk_fill_for_rank(rank, self.segments.n_segments)
                if fill is not None:
                    sheet.cell(row=target, column=segment_col).fill = fill

    # ==================================================================
    # Sheet 4 -- PSI detail
    # ==================================================================
    def _build_psi_sheet(self) -> None:
        smoothing = self.config.smoothing
        sheet, layout = self._write_detail(
            "psi", self.inputs.psi.detail,
            title="Population Stability Index - segment-level detail",
            subtitle="PSI = SUM over segments of (A - E) x ln(A / E)",
            note=(
                "A = monitoring-quarter proportion, E = reference-period proportion. All "
                f"{self.segments.n_segments} configured grades take part in every comparison, "
                "including grades with zero observations. Zero-frequency treatment: "
                f"{smoothing.describe()} PSI_CONTRIBUTION values sum to TOTAL_PSI for each "
                "model-quarter-comparison and each term is non-negative. Threshold tests use "
                "the unrounded value; displayed values are rounded for readability."
            ),
            number_formats={
                "REFERENCE_COUNT": FORMAT_COUNT, "MONITORING_COUNT": FORMAT_COUNT,
                "REFERENCE_SEGMENT_COUNT": FORMAT_COUNT, "MONITORING_SEGMENT_COUNT": FORMAT_COUNT,
                "REFERENCE_PROPORTION": FORMAT_PROPORTION,
                "MONITORING_PROPORTION": FORMAT_PROPORTION,
                "PROPORTION_CHANGE_PP": FORMAT_PP_SIGNED,
                "PSI_CONTRIBUTION": "0.000000", "TOTAL_PSI": FORMAT_INDEX,
                "CONTRIBUTION_SHARE_PCT": FORMAT_PCT,
            },
            table_name="tbl_psi_detail",
            freeze_columns=3,
        )
        if not self._can_format(layout):
            return
        apply_severity_rules(sheet, column_ref(layout, "SEVERITY"))
        apply_color_scale(sheet, column_ref(layout, "PSI_CONTRIBUTION"))
        apply_diverging_scale(sheet, column_ref(layout, "PROPORTION_CHANGE_PP"))
        apply_threshold_rules(
            sheet, column_ref(layout, "TOTAL_PSI"),
            self.config.thresholds.psi.amber, self.config.thresholds.psi.red,
        )
        apply_data_bars(sheet, column_ref(layout, "CONTRIBUTION_SHARE_PCT"))

        # PSI summary block on the same sheet, above the detail? -> separate area below
        summary = self.inputs.psi.summary
        if summary is None or summary.empty:
            return
        row = layout["last_row"] + 3
        row = write_section(
            sheet, row, "PSI SUMMARY BY MODEL, QUARTER AND COMPARISON TYPE", width=12,
            note=(
                "OVERLAPS_REFERENCE=TRUE flags a quarter that sits inside its own baseline "
                "window, where a low PSI is an artefact rather than evidence of stability."
            ),
        )
        summary_layout = write_frame(
            sheet, summary, start_row=row,
            number_formats={
                "PSI": FORMAT_INDEX, "PSI_DISPLAY": FORMAT_INDEX,
                "TOP_CONTRIBUTION": FORMAT_INDEX,
                "TOP_CONTRIBUTION_SHARE_PCT": FORMAT_PCT,
                "MAX_ABS_PROP_SHIFT_PP": FORMAT_PP_SIGNED,
                "REFERENCE_COUNT": FORMAT_COUNT, "MONITORING_COUNT": FORMAT_COUNT,
                "CONTRIBUTION_SUM_ERROR": "0.00E+00",
            },
            max_rows=self.excel.max_rows_per_sheet,
        )
        apply_severity_rules(sheet, column_ref(summary_layout, "SEVERITY"))
        apply_threshold_rules(
            sheet, column_ref(summary_layout, "PSI"),
            self.config.thresholds.psi.amber, self.config.thresholds.psi.red,
        )

    # ==================================================================
    # Sheet 5 -- CSI summary
    # ==================================================================
    def _build_csi_summary_sheet(self) -> None:
        summary = self.inputs.csi.summary
        if summary is not None and not summary.empty:
            # sort by severity then CSI, preserving every row
            ordered = summary.copy()
            ordered["__SEV"] = ordered["SEVERITY"].map(
                {"RED": 3, "AMBER": 2, "GREY": 1, "GREEN": 0}
            ).fillna(0)
            ordered = ordered.sort_values(
                ["__SEV", "CSI"], ascending=[False, False]
            ).drop(columns="__SEV")
        else:
            ordered = summary

        sheet, layout = self._write_detail(
            "csi", ordered,
            title="Characteristic Stability Index - every model, quarter and characteristic",
            subtitle="CSI_j = SUM over bins of (A_ij - E_ij) x ln(A_ij / E_ij)",
            note=(
                "Sorted by severity then CSI; no rows are removed, so the sheet remains the "
                "complete record. Bin boundaries are fitted ONCE on each model's own reference "
                "period and reused in every quarter - re-bucketing per quarter would hide the "
                "drift this metric exists to detect. Business-critical characteristics use the "
                f"tighter band (AMBER {self.config.thresholds.csi_critical.amber:g} / RED "
                f"{self.config.thresholds.csi_critical.red:g})."
            ),
            number_formats={
                "CSI": FORMAT_INDEX, "CSI_DISPLAY": FORMAT_INDEX,
                "THRESHOLD_AMBER": FORMAT_INDEX_SHORT, "THRESHOLD_RED": FORMAT_INDEX_SHORT,
                "TOP_BIN_CONTRIBUTION": FORMAT_INDEX, "TOP_BIN_SHARE_PCT": FORMAT_PCT,
                "MAX_ABS_PROP_SHIFT_PP": FORMAT_PP_SIGNED,
                "MISSING_PCT_REF": FORMAT_PCT, "MISSING_PCT_MON": FORMAT_PCT,
                "MISSINGNESS_CHANGE_PP": FORMAT_PP_SIGNED,
                "UNSEEN_CATEGORY_PCT": FORMAT_PCT, "OTHER_BIN_PCT_MON": FORMAT_PCT,
                "REFERENCE_COUNT": FORMAT_COUNT, "MONITORING_COUNT": FORMAT_COUNT,
                "UNSEEN_CATEGORY_COUNT": FORMAT_COUNT,
                "CONTRIBUTION_SUM_ERROR": "0.00E+00",
            },
            table_name="tbl_csi_summary",
            freeze_columns=4,
            width_overrides={
                "RECOMMENDED_ACTION": 80, "STATUS_REASON": 60, "BINNING_DEFINITION": 50,
                "UNSEEN_CATEGORY_EXAMPLES": 30,
            },
        )
        if not self._can_format(layout):
            return
        apply_severity_rules(sheet, column_ref(layout, "SEVERITY"))
        apply_color_scale(sheet, column_ref(layout, "CSI"))
        apply_diverging_scale(sheet, column_ref(layout, "MISSINGNESS_CHANGE_PP"))
        apply_color_scale(sheet, column_ref(layout, "UNSEEN_CATEGORY_PCT"))
        apply_data_bars(sheet, column_ref(layout, "TOP_BIN_SHARE_PCT"))

        # top-N view, without truncating the detail above
        top = self.inputs.csi.top_variables(
            n=self.config.runtime.csi_top_n or self.excel.top_n,
            quarter=self.inputs.latest_quarter,
        )
        if top.empty:
            return
        row = layout["last_row"] + 3
        row = write_section(
            sheet, row,
            f"TOP {len(top)} DRIFTING CHARACTERISTICS - {self.inputs.latest_quarter}",
            width=12,
            note="An executive extract of the table above; the full detail is retained there.",
        )
        display = top[[
            "MODEL", "VARIABLE", "VARIABLE_TYPE", "VARIABLE_ROLE", "CSI_DISPLAY", "SEVERITY",
            "TOP_CONTRIBUTING_BIN", "MISSINGNESS_CHANGE_PP", "UNSEEN_CATEGORY_PCT",
            "RECOMMENDED_ACTION",
        ]]
        top_layout = write_frame(
            sheet, display, start_row=row,
            number_formats={
                "CSI_DISPLAY": FORMAT_INDEX, "MISSINGNESS_CHANGE_PP": FORMAT_PP_SIGNED,
                "UNSEEN_CATEGORY_PCT": FORMAT_PCT,
            },
        )
        apply_severity_rules(sheet, column_ref(top_layout, "SEVERITY"))
        apply_data_bars(sheet, column_ref(top_layout, "CSI_DISPLAY"))

    # ==================================================================
    # Sheet 6 -- CSI bin detail
    # ==================================================================
    def _build_csi_bin_sheet(self) -> None:
        sheet, layout = self._write_detail(
            "csi_bins", self.inputs.csi.detail,
            title="CSI bin detail - reference vs monitoring distribution",
            subtitle="One row per model, quarter, characteristic and bin",
            note=(
                "MISSING, OTHER and UNSEEN are explicit bins, not discarded records: a "
                "characteristic whose missingness jumps is a production incident and must be "
                "visible here. Proportions are post-smoothing and sum to 1 within each "
                "characteristic; CSI_CONTRIBUTION values sum to TOTAL_CSI."
            ),
            number_formats={
                "REFERENCE_COUNT": FORMAT_COUNT, "MONITORING_COUNT": FORMAT_COUNT,
                "REFERENCE_PROPORTION": FORMAT_PROPORTION,
                "MONITORING_PROPORTION": FORMAT_PROPORTION,
                "PROPORTION_CHANGE_PP": FORMAT_PP_SIGNED,
                "CSI_CONTRIBUTION": "0.000000", "TOTAL_CSI": FORMAT_INDEX,
                "CONTRIBUTION_SHARE_PCT": FORMAT_PCT,
            },
            table_name="tbl_csi_bins",
            freeze_columns=4,
            width_overrides={"BIN": 26},
        )
        if not self._can_format(layout):
            return
        apply_severity_rules(sheet, column_ref(layout, "SEVERITY"))
        apply_color_scale(sheet, column_ref(layout, "CSI_CONTRIBUTION"))
        apply_diverging_scale(sheet, column_ref(layout, "PROPORTION_CHANGE_PP"))
        apply_data_bars(sheet, column_ref(layout, "MONITORING_PROPORTION"))

    # ==================================================================
    # Sheet 7 -- Segment migration
    # ==================================================================
    def _build_migration_sheet(self) -> None:
        sheet, layout = self._write_detail(
            "migration", self.inputs.migration_summary,
            title="Segment migration - matched customers",
            subtitle="Upgrades, downgrades, unchanged and net movement toward risk",
            note=(
                f"Computed on customers observed in BOTH periods, collapsed to one record each "
                f"using policy {self.config.runtime.multi_record_policy.value}. Upgrade = move "
                "to a safer grade under the configured risk direction. "
                + MIGRATION_CAVEAT
            ),
            number_formats={
                "FROM_CUSTOMERS": FORMAT_COUNT, "TO_CUSTOMERS": FORMAT_COUNT,
                "MATCHED_CUSTOMERS": FORMAT_COUNT, "UNCHANGED": FORMAT_COUNT,
                "UPGRADES": FORMAT_COUNT, "DOWNGRADES": FORMAT_COUNT,
                "MULTI_NOTCH_DOWNGRADES": FORMAT_COUNT, "LARGEST_FLOW_COUNT": FORMAT_COUNT,
                "LARGEST_DOWNGRADE_FLOW_COUNT": FORMAT_COUNT,
                "MATCH_RATE_PCT": FORMAT_PCT, "UNCHANGED_PCT": FORMAT_PCT,
                "UPGRADE_PCT": FORMAT_PCT, "DOWNGRADE_PCT": FORMAT_PCT,
                "MULTI_NOTCH_DOWNGRADE_PCT": FORMAT_PCT, "LARGEST_FLOW_PCT": FORMAT_PCT,
                "NET_DOWNGRADE_PP": FORMAT_PP_SIGNED,
                "AVERAGE_RANK_MOVEMENT": FORMAT_RATIO,
                "AVERAGE_ABS_RANK_MOVEMENT": FORMAT_RATIO,
            },
            table_name="tbl_migration_summary",
            freeze_columns=2,
            width_overrides={"INTERPRETATION": 100, "STATUS_REASON": 60},
            empty_message=(
                "Customer-level migration could not be produced. This needs a customer "
                "identifier present in the extract and customers observed in both periods."
            ),
        )
        if self._can_format(layout):
            apply_severity_rules(sheet, column_ref(layout, "SEVERITY"))
            apply_diverging_scale(sheet, column_ref(layout, "NET_DOWNGRADE_PP"))
            apply_color_scale(sheet, column_ref(layout, "DOWNGRADE_PCT"))
            apply_color_scale(sheet, column_ref(layout, "UPGRADE_PCT"), reverse=True)
            apply_data_bars(sheet, column_ref(layout, "MATCHED_CUSTOMERS"))

        # -- heatmap matrices for the latest quarter --------------------
        matrix = self.inputs.migration_matrix
        if matrix is None or matrix.empty:
            return
        latest = self.inputs.latest_quarter
        row = (layout["last_row"] if layout.get("rows_written") else 6) + 3
        row = write_section(
            sheet, row, f"MIGRATION MATRICES - {latest}", width=12,
            note=(
                "Row-wise percentages: each row sums to 100% of the customers who started in "
                "that grade. The diagonal is 'unchanged'; cells above it (to the right) are "
                "moves toward riskier grades."
            ),
        )
        order = list(self.segments.segments_by_risk)
        for model in self.config.active_models:
            for comparison, label in (
                (COMPARISON_REFERENCE, "vs reference period"),
                (COMPARISON_QOQ, "vs previous quarter"),
            ):
                pivot = migration_pivot(
                    matrix, model.name, comparison, latest, self.segments, value="ROW_PCT"
                )
                counts = migration_pivot(
                    matrix, model.name, comparison, latest, self.segments, value="CUSTOMER_COUNT"
                )
                if pivot.to_numpy().sum() == 0:
                    continue
                pivot = pivot.reindex(index=order, columns=order).fillna(0.0)
                counts = counts.reindex(index=order, columns=order).fillna(0)
                row = self._write_matrix_block(
                    sheet, row, f"{model.name} - {label} ({latest})", pivot, counts
                )
                row += 1

    def _write_matrix_block(
        self,
        sheet: Worksheet,
        row: int,
        title: str,
        pivot: pd.DataFrame,
        counts: pd.DataFrame,
    ) -> int:
        """Write one 6x6 matrix (row %) with a count matrix beside it."""
        cell = sheet.cell(row=row, column=1, value=title)
        cell.font = Font(bold=True, size=10, color=NAVY)
        row += 1
        n_segments = self.segments.n_segments

        header_row = row
        sheet.cell(row=header_row, column=1, value="FROM \\ TO (row %)").font = Font(
            bold=True, size=9
        )
        for offset, segment in enumerate(pivot.columns, start=2):
            header = sheet.cell(row=header_row, column=offset, value=str(segment))
            header.font = Font(bold=True, size=9)
            fill = risk_fill_for_rank(self.segments.risk_rank.get(str(segment)), n_segments)
            if fill is not None:
                header.fill = fill
            header.alignment = Alignment(horizontal="center")
        count_offset = len(pivot.columns) + 3
        sheet.cell(
            row=header_row, column=count_offset, value="FROM \\ TO (customers)"
        ).font = Font(bold=True, size=9)
        for offset, segment in enumerate(counts.columns, start=count_offset + 1):
            header = sheet.cell(row=header_row, column=offset, value=str(segment))
            header.font = Font(bold=True, size=9)
            header.alignment = Alignment(horizontal="center")

        for row_offset, segment in enumerate(pivot.index, start=1):
            target = header_row + row_offset
            label = sheet.cell(row=target, column=1, value=str(segment))
            label.font = Font(bold=True, size=9)
            fill = risk_fill_for_rank(self.segments.risk_rank.get(str(segment)), n_segments)
            if fill is not None:
                label.fill = fill
            for col_offset, column in enumerate(pivot.columns, start=2):
                value = float(pivot.loc[segment, column])
                target_cell = sheet.cell(row=target, column=col_offset, value=value)
                target_cell.number_format = FORMAT_PCT
                if str(segment) == str(column):
                    target_cell.font = Font(bold=True, size=9)
            sheet.cell(row=target, column=count_offset, value=str(segment)).font = Font(
                bold=True, size=9
            )
            for col_offset, column in enumerate(counts.columns, start=count_offset + 1):
                target_cell = sheet.cell(
                    row=target, column=col_offset, value=int(counts.loc[segment, column])
                )
                target_cell.number_format = FORMAT_COUNT

        last_row = header_row + len(pivot.index)
        percent_ref = (
            f"{get_column_letter(2)}{header_row + 1}:"
            f"{get_column_letter(1 + len(pivot.columns))}{last_row}"
        )
        count_ref = (
            f"{get_column_letter(count_offset + 1)}{header_row + 1}:"
            f"{get_column_letter(count_offset + len(counts.columns))}{last_row}"
        )
        apply_heatmap(sheet, percent_ref)
        apply_heatmap(sheet, count_ref)
        return last_row + 2

    # ==================================================================
    # Sheet 8 -- Data quality
    # ==================================================================
    def _build_data_quality_sheet(self) -> None:
        sheet, layout = self._write_detail(
            "data_quality", self.inputs.data_quality,
            title="Data quality checks and exceptions",
            subtitle="Missingness, duplication, date integrity, outliers, coverage and freshness",
            note=(
                "Sorted most severe first. GREY marks a check that could not be concluded "
                "rather than one that passed. Every row names the affected population and the "
                "recommended remediation; nothing here is auto-corrected in the data."
            ),
            number_formats={"VALUE": "0.00", "RECORDS_AFFECTED": FORMAT_COUNT},
            table_name="tbl_data_quality",
            freeze_columns=3,
            width_overrides={"DETAIL": 90, "RECOMMENDED_ACTION": 90, "THRESHOLD": 34},
        )
        if not self._can_format(layout):
            return
        apply_severity_rules(sheet, column_ref(layout, "SEVERITY"))
        apply_data_bars(sheet, column_ref(layout, "RECORDS_AFFECTED"))

    # ==================================================================
    # Sheet 9 -- Outcome performance
    # ==================================================================
    def _build_performance_sheet(self) -> None:
        performance = self.inputs.performance
        sheet = self._sheet("performance")
        width = 12
        outcome = self.config.outcome

        row = write_title(
            sheet,
            "Outcome-based predictive performance",
            "Bad rates, discrimination, lift, rank ordering, calibration and vintages",
            width=width,
            note=(
                "Model STABILITY (PSI/CSI/migration) and model PREDICTIVE PERFORMANCE are "
                "different questions. A stable population does not imply the scorecard still "
                "discriminates, and a population shift does not by itself prove it does not."
            ),
        )

        if not performance.enabled or performance.summary.empty:
            row = write_section(
                sheet, row, "STATUS: NOT AVAILABLE", width=width,
                note=(
                    "No usable outcome label was supplied, so every performance metric below is "
                    "reported as NOT AVAILABLE with its missing dependency named. The pipeline "
                    "does not estimate discrimination from the segment distribution: a stable "
                    "population tells you nothing about whether the model still separates good "
                    "from bad."
                ),
            )
            row = write_key_values(sheet, row, [
                ("Outcome column configured", outcome.outcome_column or "NONE"),
                ("Observation window (months)", outcome.observation_window_months),
                ("Expected PD column", outcome.pd_column or "NONE"),
                ("Vintage date column", outcome.vintage_date_column or self.config.date_column),
                ("Immature accounts excluded", str(outcome.exclude_immature)),
            ], value_col=4, label_width=3)
            row += 1
            row = write_section(sheet, row, "METRIC AVAILABILITY", width=width)
            layout = write_frame(sheet, performance.availability, start_row=row)
            apply_severity_rules(sheet, column_ref(layout, "STATUS"))
            autosize_columns(
                sheet, performance.availability,
                overrides={"EXPLANATION": 100, "MISSING_DEPENDENCY": 40, "METRIC": 42},
            )
            finalise(sheet, layout, freeze=True, autofilter=True)
            return

        row = write_section(
            sheet, row, "OUTCOME DEFINITION", width=width,
            note=(
                "Every figure on this sheet is relative to this definition. A different bad "
                "definition, window or population gives different numbers for the same book."
            ),
        )
        definition_items = [
            ("Definition", outcome.describe()),
            ("Outcome column", outcome.outcome_column or NOT_AVAILABLE),
            ("Basis", outcome.definition_basis.value),
            ("Supplied at", f"{outcome.outcome_level.value} level"),
            ("Metrics measured on", f"one record per {outcome.metrics_level.value.lower()}"),
            ("Bad values", ", ".join(str(v) for v in outcome.bad_values)),
            ("Good values", ", ".join(str(v) for v in outcome.good_values)),
            ("Indeterminate values (excluded, NOT good)",
             ", ".join(str(v) for v in outcome.indeterminate_values) or "none declared"),
            ("Observation window (months)", outcome.observation_window_months),
            ("Outcome as-of date",
             outcome.outcome_as_of.isoformat() if outcome.outcome_as_of else "max(login date)"),
            ("Expected PD column", outcome.pd_column or "none - calibration NOT AVAILABLE"),
        ]
        row = write_key_values(
            sheet, row, definition_items, value_col=4, label_width=3,
            severities={"Indeterminate values (excluded, NOT good)":
                        "GREEN" if outcome.indeterminate_values else "AMBER"},
        )
        caveat = outcome.proxy_caveat()
        if caveat:
            row = write_note(sheet, row, caveat, width=width) + 1
        else:
            row += 1

        row = write_section(
            sheet, row, "PERFORMANCE SUMMARY BY MODEL AND QUARTER", width=width,
            note=(
                f"Performance window: {outcome.observation_window_months} months. Accounts whose "
                "window has not elapsed are EXCLUDED rather than counted as good - an account "
                "with no default yet is an unknown, not a good customer. COHORT_COMPLETE=FALSE "
                "means the bad rate shown is a censored lower bound that will still rise."
            ),
        )
        layout = write_frame(
            sheet, performance.summary, start_row=row,
            number_formats={
                "TOTAL_RECORDS": FORMAT_COUNT, "RECORDS_WITH_OUTCOME": FORMAT_COUNT,
                "MATURED_RECORDS": FORMAT_COUNT, "IMMATURE_EXCLUDED": FORMAT_COUNT,
                "INDETERMINATE_EXCLUDED": FORMAT_COUNT, "UNKNOWN_EXCLUDED": FORMAT_COUNT,
                "APPLICATIONS_BEFORE_DEDUPE": FORMAT_COUNT, "UNITS_IN_METRIC": FORMAT_COUNT,
                "BAD_COUNT": FORMAT_COUNT, "GOOD_COUNT": FORMAT_COUNT,
                "MATURITY_PCT": FORMAT_PCT, "EVENT_RATE_PCT": FORMAT_PCT,
                "NON_EVENT_RATE_PCT": FORMAT_PCT, "EXPECTED_BAD_RATE_PCT": FORMAT_PCT,
                "AUC": FORMAT_INDEX, "GINI": FORMAT_INDEX, "KS_STATISTIC": FORMAT_INDEX,
                "LIFT_TOP_SEGMENT": FORMAT_RATIO,
                "BAD_CAPTURE_TOP_2_SEGMENTS_PCT": FORMAT_PCT,
                "OBSERVED_VS_EXPECTED_RATIO": FORMAT_RATIO, "CALIBRATION_Z": FORMAT_RATIO,
                "SPEARMAN_RANK_CORRELATION": FORMAT_RATIO,
            },
            max_rows=self.excel.max_rows_per_sheet,
        )
        apply_severity_rules(sheet, column_ref(layout, "PERFORMANCE_SEVERITY"))
        apply_color_scale(sheet, column_ref(layout, "GINI"), reverse=True)
        apply_color_scale(sheet, column_ref(layout, "KS_STATISTIC"), reverse=True)
        apply_color_scale(sheet, column_ref(layout, "EVENT_RATE_PCT"))
        apply_data_bars(sheet, column_ref(layout, "MATURITY_PCT"))
        autosize_columns(
            sheet, performance.summary,
            overrides={
                "PERFORMANCE_STATUS": 100, "RANK_ORDERING": 54, "CALIBRATION_VERDICT": 60,
                "SCORE_SOURCE": 40, "OUTCOME_DEFINITION": 60,
            },
        )
        finalise(
            sheet, layout, freeze=True, autofilter=True,
            table_name="tbl_performance" if self.excel.use_excel_tables else None,
            table_row_limit=self.excel.excel_table_row_limit, freeze_columns=2,
        )
        row = layout["last_row"] + 3

        row = write_section(
            sheet, row, "BAD RATE, LIFT AND CUMULATIVE BAD CAPTURE BY SEGMENT", width=width,
            note=(
                "Ordered riskiest grade first, so a well-ordered scorecard shows a falling bad "
                "rate down the rows and a cumulative bad capture that rises faster than the "
                "cumulative population. RANK_ORDER_BREAK marks a grade whose bad rate is higher "
                "than the riskier grade above it."
            ),
        )
        segment_layout = write_frame(
            sheet, performance.by_segment, start_row=row,
            number_formats={
                "MATURED_RECORDS": FORMAT_COUNT, "BAD_COUNT": FORMAT_COUNT,
                "BAD_RATE_PCT": FORMAT_PCT, "SHARE_OF_POPULATION_PCT": FORMAT_PCT,
                "SHARE_OF_BADS_PCT": FORMAT_PCT, "LIFT": FORMAT_RATIO,
                "CUMULATIVE_POPULATION_PCT": FORMAT_PCT,
                "CUMULATIVE_BAD_CAPTURE_PCT": FORMAT_PCT,
                "EXPECTED_PD_PCT": FORMAT_PCT, "OBSERVED_VS_EXPECTED_RATIO": FORMAT_RATIO,
            },
            max_rows=self.excel.max_rows_per_sheet,
        )
        apply_color_scale(sheet, column_ref(segment_layout, "BAD_RATE_PCT"))
        apply_data_bars(sheet, column_ref(segment_layout, "CUMULATIVE_BAD_CAPTURE_PCT"))
        apply_color_scale(sheet, column_ref(segment_layout, "LIFT"))
        row = segment_layout["last_row"] + 3

        if not performance.vintage.empty:
            row = write_section(
                sheet, row, "VINTAGE ANALYSIS", width=width,
                note=(
                    "Bad rate by origination quarter and grade. Cohorts flagged "
                    "COHORT_COMPLETE=FALSE are right-censored and are NOT comparable with "
                    "seasoned vintages, however attractive their bad rate looks."
                ),
            )
            vintage_layout = write_frame(
                sheet, performance.vintage, start_row=row,
                number_formats={
                    "ACCOUNTS": FORMAT_COUNT, "MATURED_ACCOUNTS": FORMAT_COUNT,
                    "BAD_COUNT": FORMAT_COUNT, "MATURITY_PCT": FORMAT_PCT,
                    "BAD_RATE_PCT": FORMAT_PCT,
                },
                max_rows=self.excel.max_rows_per_sheet,
            )
            apply_color_scale(sheet, column_ref(vintage_layout, "BAD_RATE_PCT"))
            apply_data_bars(sheet, column_ref(vintage_layout, "MATURITY_PCT"))
            row = vintage_layout["last_row"] + 3

        row = write_section(sheet, row, "METRIC AVAILABILITY", width=width)
        availability_layout = write_frame(sheet, performance.availability, start_row=row)
        apply_severity_rules(sheet, column_ref(availability_layout, "STATUS"))

    # ==================================================================
    # Sheet 10 -- Configuration and methodology
    # ==================================================================
    def _build_config_sheet(self) -> None:
        from .methodology import methodology_blocks

        sheet = self._sheet("config")
        width = 10
        row = write_title(
            sheet,
            "Configuration, methodology and limitations",
            "Everything needed to reproduce and challenge this pack",
            width=width,
            note=(
                "Governance evidence: column mappings, reference periods and modes, formulas, "
                "the zero-frequency treatment, thresholds, binning rules and the documented "
                "assumptions and limitations."
            ),
        )
        for title, note, frame, formats in methodology_blocks(self.inputs):
            row = write_section(sheet, row, title, width=width, note=note)
            if frame is None or frame.empty:
                row = write_note(sheet, row, "No entries.", width=width) + 1
                continue
            layout = write_frame(
                sheet, frame, start_row=row, number_formats=formats,
                max_rows=self.excel.max_rows_per_sheet, header_wrap=False,
            )
            for name in ("SEVERITY", "STATUS", "SUFFICIENT_SAMPLE"):
                apply_severity_rules(sheet, column_ref(layout, name))
            row = layout["last_row"] + 2
        sheet.column_dimensions["A"].width = 46
        sheet.column_dimensions["B"].width = 58
        for letter in ("C", "D", "E", "F", "G", "H", "I", "J"):
            sheet.column_dimensions[letter].width = 22
        sheet.sheet_view.showGridLines = False

    # ==================================================================
    # Sheet 11 -- Alert register
    # ==================================================================
    def _build_alert_sheet(self) -> None:
        sheet, layout = self._write_detail(
            "alerts", self.inputs.alerts,
            title="Alert register - unified monitoring alerts",
            subtitle="GREEN within limits | AMBER investigate | RED escalate | GREY insufficient data",
            note=(
                "One row per evaluated metric, sorted most severe first. MANUAL_REVIEW_STATUS "
                "and MODEL_OWNER_NOTES are left blank for the model owner to complete and sign "
                "off. A PSI or CSI breach triggers investigation under the model risk policy; "
                "it does not by itself invalidate the scorecard."
            ),
            number_formats={"METRIC_VALUE": "0.0000"},
            table_name="tbl_alerts",
            freeze_columns=4,
            width_overrides={
                "EXPLANATION": 100, "RECOMMENDED_ACTION": 100, "THRESHOLD": 34,
                "METRIC": 34, "VARIABLE_OR_SEGMENT": 28, "MODEL_OWNER_NOTES": 36,
                "INVESTIGATION_STATUS": 32, "MANUAL_REVIEW_STATUS": 22,
            },
        )
        if not self._can_format(layout):
            return
        apply_severity_rules(sheet, column_ref(layout, "SEVERITY"))
        notes_column = layout["columns"].get("MODEL_OWNER_NOTES")
        if notes_column:
            highlight = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
            for target in range(layout["first_data_row"], layout["last_row"] + 1):
                sheet.cell(row=target, column=notes_column).fill = highlight

    # ==================================================================
    # Sheet 12 -- Execution log
    # ==================================================================
    def _build_log_sheet(self) -> None:
        log_frame = pd.DataFrame(
            self.inputs.execution_log.to_rows(),
            columns=["TIMESTAMP", "LEVEL", "STAGE", "MESSAGE", "DETAIL"],
        )
        sheet, layout = self._write_detail(
            "log", log_frame,
            title="Execution log - warnings and exceptions",
            subtitle=(
                f"{self.inputs.execution_log.error_count} error(s), "
                f"{self.inputs.execution_log.warning_count} warning(s) during this run"
            ),
            note=(
                "Every warning raised while the pack was produced, so a reader can see what the "
                "pipeline could not do as well as what it did. An empty log means no exception "
                "was raised, not that no checks ran."
            ),
            table_name="tbl_execution_log",
            freeze_columns=2,
            width_overrides={"MESSAGE": 120, "DETAIL": 60},
            empty_message="No warnings or errors were raised during this execution.",
        )
        if not self._can_format(layout):
            return
        for label, colour in (("ERROR", "RED"), ("WARNING", "AMBER"), ("INFO", "GREEN")):
            reference = column_ref(layout, "LEVEL")
            if reference:
                from openpyxl.formatting.rule import CellIsRule

                sheet.conditional_formatting.add(
                    reference,
                    CellIsRule(
                        operator="equal", formula=[f'"{label}"'],
                        fill=severity_fill(colour), font=severity_font(colour),
                    ),
                )

    # ==================================================================
    # Sheet 13 -- Bin definitions
    # ==================================================================
    def _build_bin_definition_sheet(self) -> None:
        sheet, layout = self._write_detail(
            "bins", self.inputs.csi.definitions,
            title="Fitted reference bin definitions",
            subtitle="The exact boundaries used for every CSI comparison in this pack",
            note=(
                "Fitted once on each model's own reference period and reused unchanged across "
                "all monitoring quarters. Retaining them is what makes a CSI figure auditable "
                "and reproducible: BINNING_SOURCE states whether approved scorecard boundaries, "
                "reference quantiles or reference categories were used."
            ),
            number_formats={
                "REFERENCE_RECORDS": FORMAT_COUNT, "REFERENCE_MISSING": FORMAT_COUNT,
                "REFERENCE_DISTINCT_VALUES": FORMAT_COUNT,
            },
            table_name="tbl_bin_definitions",
            freeze_columns=2,
            width_overrides={"DEFINITION": 90, "NOTES": 80},
            empty_message="No bin definitions were fitted (no characteristics were monitored).",
        )
        if self._can_format(layout):
            apply_severity_rules(sheet, column_ref(layout, "USABLE"))

    # ==================================================================
    # Sheet 14 -- Test scenarios
    # ==================================================================
    def _build_scenario_sheet(self) -> None:
        self._write_detail(
            "scenarios", self.inputs.scenarios,
            title="Planted test scenarios (synthetic dataset)",
            subtitle="What this run is expected to surface, and where",
            note=(
                "The synthetic generator plants each condition deliberately so that a run "
                "exercises the real code paths. Use this sheet to verify the pipeline is "
                "detecting what it claims to detect."
            ),
            table_name="tbl_scenarios",
            freeze_columns=1,
            width_overrides={"SCENARIO": 50, "WHERE": 50, "EXPECT": 90},
        )



def build_workbook(inputs: ReportInputs, path: str | Path) -> Path:
    """Convenience wrapper around :class:`ExcelReportBuilder`."""
    return ExcelReportBuilder(inputs).build(path)


__all__ = ["ReportInputs", "ExcelReportBuilder", "build_workbook", "SHEET_NAMES"]
